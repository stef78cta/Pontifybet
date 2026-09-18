"""Evaluator de audit pentru formulele nucleului acestui workbook, nu motor Excel general.

Rulează pe cele cinci rânduri DEMO și compară cu valorile cache din atașament.
Nu certifică Microsoft Excel și nu implementează funcțiile SUMIFS/COUNTIFS de audit.
"""
import json, math, re, sys
from pathlib import Path
from functools import lru_cache
from decimal import Decimal, ROUND_HALF_UP
import openpyxl
from openpyxl.formula.tokenizer import Tokenizer
from openpyxl.utils.cell import range_boundaries, get_column_letter

ROOT=Path(__file__).resolve().parent.parent
SOURCE=ROOT/'upload/1_model_analiza_over0.5_optimizat_V4.xlsx'
class XLException(Exception): pass

@lru_cache(None)
def parse(formula):
    ts=[t for t in Tokenizer(formula).items if t.type!='WHITE-SPACE']; i=0
    prec={'=':1,'<>':1,'<':1,'>':1,'<=':1,'>=':1,'&':2,'+':3,'-':3,'*':4,'/':4,'^':5}
    def expr(p=0):
        nonlocal i
        t=ts[i];i+=1
        if t.type=='OPERAND':
            a=('ref',t.value) if t.subtype=='RANGE' else ('lit',t.value[1:-1].replace('""','"') if t.subtype=='TEXT' else float(t.value) if t.subtype=='NUMBER' else t.value=='TRUE')
        elif t.type=='OPERATOR-PREFIX': a=('unary',t.value,expr(6))
        elif t.type=='PAREN' and t.subtype=='OPEN':
            a=expr();assert ts[i].subtype=='CLOSE';i+=1
        elif t.type=='FUNC' and t.subtype=='OPEN':
            args=[]
            while not (ts[i].type=='FUNC' and ts[i].subtype=='CLOSE'):
                args.append(expr())
                if ts[i].type=='SEP':i+=1
                else:break
            assert ts[i].subtype=='CLOSE';i+=1;a=('call',t.value[:-1],args)
        else: raise ValueError((t.value,formula))
        while i<len(ts):
            t=ts[i]
            if t.type=='OPERATOR-POSTFIX':i+=1;a=('op','/',a,('lit',100));continue
            if t.type!='OPERATOR-INFIX' or prec[t.value]<p:break
            i+=1;a=('op',t.value,a,expr(prec[t.value]+1))
        return a
    ast=expr(); assert i==len(ts),(formula,i,len(ts));return ast

def flatten(v):
    if isinstance(v,list):
        for x in v:yield from flatten(x)
    else:yield v
def isnum(v):return isinstance(v,(int,float)) and not isinstance(v,bool)
def num(v):
    if v is None or v=='':return 0.0
    if isinstance(v,(int,float)):return v
    raise XLException('VALUE')
def txt(v):
    if v is None:return ''
    if isnum(v):return str(int(v)) if v==int(v) else str(v)
    return str(v)
def compare(a,b,op):
    if a is None:a='' if isinstance(b,str) else 0
    if b is None:b='' if isinstance(a,str) else 0
    # Ordine Excel pentru valori scalare: numeric înainte de text.
    ka=(1,a.casefold()) if isinstance(a,str) else (0,a)
    kb=(1,b.casefold()) if isinstance(b,str) else (0,b)
    return {'=':lambda:ka==kb,'<>':lambda:ka!=kb,'<':lambda:ka<kb,'>':lambda:ka>kb,'<=':lambda:ka<=kb,'>=':lambda:ka>=kb}[op]()

class AuditEngine:
    def __init__(self,wb,overrides=None):self.wb=wb;self.overrides=overrides or {};self.memo={};self.visiting=set()
    def cell(self,s,c):
        c=c.replace('$','');key=(s,c)
        if key in self.overrides:return self.overrides[key]
        if key in self.memo:return self.memo[key]
        if key in self.visiting:raise XLException('CYCLE')
        raw=self.wb[s][c].value
        if not (isinstance(raw,str) and raw.startswith('=')):return raw
        self.visiting.add(key)
        try:r=self.ev(parse(raw),s)
        finally:self.visiting.remove(key)
        self.memo[key]=r;return r
    def ref(self,ref,s):
        if '!' in ref:s,ref=ref.rsplit('!',1);s=s.strip("'").replace("''", "'")
        ref=ref.replace('$','')
        if ':' not in ref:return self.cell(s,ref)
        a,b,c,d=range_boundaries(ref)
        return [self.cell(s,f'{get_column_letter(x)}{y}') for y in range(b,d+1) for x in range(a,c+1)]
    def ev(self,a,s):
        typ=a[0]
        if typ=='lit':return a[1]
        if typ=='ref':return self.ref(a[1],s)
        if typ=='unary':return num(self.ev(a[2],s))*(1 if a[1]=='+' else -1)
        if typ=='op':
            op=a[1];x=self.ev(a[2],s);y=self.ev(a[3],s)
            if op in ['=','<>','<','>','<=','>=']:return compare(x,y,op)
            if op=='&':return txt(x)+txt(y)
            x=num(x);y=num(y)
            return {'+':lambda:x+y,'-':lambda:x-y,'*':lambda:x*y,'/':lambda:x/y,'^':lambda:x**y}[op]()
        name,args=a[1],a[2]
        if name=='IF':return self.ev(args[1] if self.ev(args[0],s) else args[2],s)
        if name=='IFERROR':
            try:return self.ev(args[0],s)
            except (XLException,ZeroDivisionError,ValueError,OverflowError):return self.ev(args[1],s)
        v=[self.ev(x,s) for x in args];ns=[x for x in flatten(v) if isnum(x)]
        if name=='COUNT':return len(ns)
        if name=='SUM':return sum(ns)
        if name=='AVERAGE':return sum(ns)/len(ns)
        if name=='MIN':return min(ns) if ns else 0
        if name=='MAX':return max(ns) if ns else 0
        if name=='AND':return all(v)
        if name=='OR':return any(v)
        if name=='ABS':return abs(num(v[0]))
        if name=='SQRT':return math.sqrt(num(v[0]))
        if name=='EXP':return math.exp(num(v[0]))
        if name=='LN':return math.log(num(v[0]))
        if name=='POWER':return num(v[0])**num(v[1])
        if name=='ROUND':return float(Decimal(str(v[0])).quantize(Decimal(1).scaleb(-int(v[1])),rounding=ROUND_HALF_UP))
        if name=='MATCH':
            for i,x in enumerate(v[1],1):
                if compare(v[0],x,'='):return i
            raise XLException('NA')
        if name=='INDEX':return v[0][int(v[1])-1]
        if name=='TEXT':
            fmt=v[1];digits=len(fmt.split('.')[1].replace('%','').replace('x','')) if '.' in fmt else 0
            z=num(v[0])*(100 if '%' in fmt else 1)
            z=Decimal(str(z)).quantize(Decimal(1).scaleb(-digits),rounding=ROUND_HALF_UP)
            return f'{z:.{digits}f}'+('%' if '%' in fmt else 'x' if fmt.endswith('x') else '')
        raise NotImplementedError(name)

def main():
    w=openpyxl.load_workbook(SOURCE);cached=openpyxl.load_workbook(SOURCE,data_only=True)
    eng=AuditEngine(w);comparisons=[];fixtures=[]
    for r in range(4,9):
        s=w['Analize meciuri'];outputs={};inputs={}
        for c in s[r]:
            if c.data_type!='f':inputs[c.column_letter]=c.value;continue
            actual=eng.cell(s.title,c.coordinate);expected=cached[s.title][c.coordinate].value
            ok=math.isclose(actual,expected,rel_tol=1e-12,abs_tol=1e-12) if isnum(actual) and isnum(expected) else actual==expected or actual=='' and expected is None
            comparisons.append({'cell':c.coordinate,'ok':ok,'cache':expected,'calculated':actual,'abs_difference':abs(actual-expected) if isnum(actual) and isnum(expected) else None})
            outputs[c.column_letter]={'cache':expected,'simulated':actual}
        fixtures.append({'match_id':s[f'A{r}'].value,'excel_row':r,'inputs':inputs,'outputs':outputs})
    scenarios={
      'missing_EV':{'EV':None},'missing_xG_L':{'L':None},'missing_recent_blank_FY':{'FY':None},
      'small_sample':{'I':4,'P':4,'W':3,'AB':3},'no_odds':{'AR':None,'AS':None},
      'confirmed_lineup':{'AU':1},'striker_absent':{'CY':0},'unused_FB':{'FB':2},
      'odds_changed':{'AR':1.5,'AS':2.5},'missing_O05':{'N':None,'U':None},
      'low_conversion':{'J':1,'K':1,'L':1,'M':1,'Q':1,'R':1,'S':1,'T':1,'X':0.8,'Y':1,'Z':1,'AA':1,'AC':0.8,'AD':1,'AE':1,'AF':1,'AJ':1,'AK':1,'AL':2,'AH':2,'AV':0,'AW':0,'AY':0,'BF':0},
      'context_text':{'AV':'NOT CHECKED'},'context_blank':{'AV':None},
    }
    results=[]
    for name,ov in scenarios.items():
        e=AuditEngine(w,{('Analize meciuri',k+'4'):v for k,v in ov.items()});o={}
        for col in ['GJ','GM','GL','GO','GP','GT','HE','HG','HH','HK']:
            try:o[col]=e.cell('Analize meciuri',col+'4')
            except Exception as ex:o[col]='ERROR:'+type(ex).__name__
        results.append({'name':name,'overrides':ov,'simulated':o,'native_excel_verified':False})
    out={'scope':'Independent formula evaluation; saved cache is not a fresh Microsoft Excel run. TEXT emulation uses dot decimal.','tolerance':1e-12,'comparisons':comparisons,'fixtures':fixtures,'scenarios':results}
    (ROOT/'analysis/verification.json').write_text(json.dumps(out,ensure_ascii=False,indent=2,default=str))
    failures=[x for x in comparisons if not x['ok']]
    print('Formula comparisons',len(comparisons),'PASS',len(comparisons)-len(failures),'FAIL',len(failures))
    print(json.dumps(failures,ensure_ascii=False,indent=2));print(json.dumps(results,ensure_ascii=False,indent=2))
if __name__=='__main__':main()
