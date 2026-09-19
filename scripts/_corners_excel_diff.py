"""Diferențele dintre cache-ul V14 și recalcularea Microsoft Excel (template intact)."""

from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from config.settings import TEMPLATES_DIR  # noqa: E402
from src.engines.corners.engine import TEMPLATE_NAME  # noqa: E402

TEMPLATE = TEMPLATES_DIR / TEMPLATE_NAME
RECALC = ROOT / ".probe" / "pristine.xlsx"

SHEETS = {
    "Motor_Meciuri": range(6, 9),
    "Core_Nativ": range(6, 42),
    "Distributie": range(6, 9),
    "Analiza_Linii": range(6, 48),
    "Verificari": range(1, 16),
}


def same(a, b) -> bool:
    if a is None and b is None:
        return True
    if isinstance(a, (int, float)) and isinstance(b, (int, float)):
        return abs(float(a) - float(b)) <= max(1e-9, 1e-9 * abs(float(b)))
    return a == b


def main() -> None:
    cached = load_workbook(TEMPLATE, data_only=True)
    recalc = load_workbook(RECALC, data_only=True)
    formulas = load_workbook(TEMPLATE, data_only=False)
    total = 0
    for sheet, rows in SHEETS.items():
        ws_c, ws_r, ws_f = cached[sheet], recalc[sheet], formulas[sheet]
        shown = 0
        for row in rows:
            for col in range(1, ws_f.max_column + 1):
                cell = ws_f.cell(row=row, column=col)
                if not isinstance(cell.value, str) or not cell.value.startswith("="):
                    continue
                a = ws_c.cell(row=row, column=col).value
                b = ws_r.cell(row=row, column=col).value
                if same(a, b):
                    continue
                total += 1
                if shown < 12:
                    shown += 1
                    print(f"{sheet}!{cell.coordinate}: cache={a!r} excel={b!r}")
                    print(f"    {cell.value[:220]}")
        print(f"--- {sheet}: {shown} afișate")
    print(f"TOTAL diferențe: {total}")


if __name__ == "__main__":
    main()
