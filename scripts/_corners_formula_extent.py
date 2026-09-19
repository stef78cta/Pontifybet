"""Raportează extinderea reală a formulelor în șablonul Cornere V14."""

from __future__ import annotations

from openpyxl import load_workbook

from config.settings import TEMPLATES_DIR

TEMPLATE = TEMPLATES_DIR / "6_model_analiza_cornere_multiline_optimizat_V14.xlsx"


def main() -> None:
    wb = load_workbook(TEMPLATE, data_only=False)
    total = 0
    print(f"{'sheet':22} {'formule':>8} {'max_row':>8} {'max_col':>8}")
    for name in wb.sheetnames:
        ws = wb[name]
        count = 0
        max_row = 0
        max_col = 0
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    count += 1
                    max_row = max(max_row, cell.row)
                    max_col = max(max_col, cell.column)
        total += count
        print(f"{name:22} {count:8d} {max_row:8d} {max_col:8d}")
    print(f"TOTAL formule: {total}")
    wb.close()


if __name__ == "__main__":
    main()
