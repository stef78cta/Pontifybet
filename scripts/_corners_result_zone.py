"""Verifică dacă zona de rezultate propusă în `Input_Meci` este liberă.

Liberă înseamnă: fără valori, fără formule proprii și fără referințe din alte foi
către rândurile respective. Rulează înainte de a scrie rezultatele Python în export.
"""

from __future__ import annotations

import re

from openpyxl import load_workbook

from config.settings import TEMPLATES_DIR

TEMPLATE = TEMPLATES_DIR / "6_model_analiza_cornere_multiline_optimizat_V14.xlsx"
SHEET = "Input_Meci"
FIRST_ROW = 107
LAST_ROW = 140

REF = re.compile(r"Input_Meci!\$?([A-Z]{1,3})\$?(\d+)")


def main() -> None:
    wb = load_workbook(TEMPLATE, data_only=False)
    ws = wb[SHEET]
    print(f"{SHEET}: dimensiuni={ws.dimensions} max_row={ws.max_row} max_col={ws.max_column}")
    print(f"tabele: {list(ws.tables)}")
    for name in ws.tables:
        print(f"  tabel {name}: {ws.tables[name].ref}")

    occupied: list[str] = []
    for row in ws.iter_rows(min_row=FIRST_ROW, max_row=max(LAST_ROW, ws.max_row)):
        for cell in row:
            if cell.value is not None:
                occupied.append(f"{cell.coordinate}={cell.value!r}")
    print(f"celule ocupate de la rândul {FIRST_ROW}: {len(occupied)}")
    for item in occupied[:20]:
        print("   ", item)

    max_ref_row = 0
    samples: list[str] = []
    for sheet_name in wb.sheetnames:
        sheet = wb[sheet_name]
        for row in sheet.iter_rows():
            for cell in row:
                if not isinstance(cell.value, str) or "Input_Meci" not in cell.value:
                    continue
                for _, digits in REF.findall(cell.value):
                    number = int(digits)
                    if number > max_ref_row:
                        max_ref_row = number
                        samples = [f"{sheet_name}!{cell.coordinate}: {cell.value[:120]}"]
                    elif number == max_ref_row and len(samples) < 3:
                        samples.append(f"{sheet_name}!{cell.coordinate}: {cell.value[:120]}")
    print(f"cel mai mare rând Input_Meci referit de formule: {max_ref_row}")
    for item in samples:
        print("   ", item)
    wb.close()


if __name__ == "__main__":
    main()
