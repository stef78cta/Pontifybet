"""Probe: ce se întâmplă cu formulele V14 la un round-trip openpyxl + recalc Excel."""

from __future__ import annotations

import shutil
import sys
import time
import zipfile
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from config.settings import TEMPLATES_DIR  # noqa: E402
from src.engines.corners.engine import TEMPLATE_NAME  # noqa: E402

TEMPLATE = TEMPLATES_DIR / TEMPLATE_NAME
OUT = ROOT / ".probe"
OUT.mkdir(exist_ok=True)


def zip_parts(path: Path) -> list[str]:
    with zipfile.ZipFile(path) as zf:
        return sorted(zf.namelist())


def dynamic_array_cells(path: Path) -> int:
    """Numără celulele marcate cu metadate de dynamic array (`cm=`)."""
    count = 0
    with zipfile.ZipFile(path) as zf:
        for name in zf.namelist():
            if not name.startswith("xl/worksheets/sheet"):
                continue
            data = zf.read(name).decode("utf-8", "replace")
            count += data.count(' cm="')
    return count


def recalc(path: Path) -> None:
    import pythoncom
    import win32com.client

    pythoncom.CoInitialize()
    excel = win32com.client.DispatchEx("Excel.Application")
    excel.Visible = False
    excel.DisplayAlerts = False
    excel.AskToUpdateLinks = False
    wb = None
    try:
        wb = excel.Workbooks.Open(str(path), UpdateLinks=0, ReadOnly=False)
        excel.Calculation = -4105
        excel.CalculateFullRebuild()
        deadline = time.time() + 600
        while int(excel.CalculationState) != 0 and time.time() < deadline:
            time.sleep(0.5)
        print(f"  CalculationState={int(excel.CalculationState)}")
        wb.Save()
        wb.Close(SaveChanges=True)
        wb = None
    finally:
        if wb is not None:
            wb.Close(SaveChanges=False)
        excel.Quit()
        pythoncom.CoUninitialize()


def report(path: Path, label: str) -> None:
    wb = load_workbook(path, data_only=True)
    ws = wb["Dashboard"]
    dash = {letter: ws[f"{letter}26"].value for letter in ("E", "G", "H", "I", "K")}
    motor = wb["Motor_Meciuri"]
    core = wb["Core_Nativ"]
    print(f"[{label}] dynamic_array_cells={dynamic_array_cells(path)}")
    print(f"[{label}] Dashboard26={dash}")
    print(f"[{label}] Motor_Meciuri A6={motor['A6'].value!r} H6={motor['H6'].value!r}")
    print(f"[{label}] Core_Nativ A6={core['A6'].value!r} E6={core['E6'].value!r}")
    wb.close()


def main() -> None:
    print(f"template dynamic_array_cells={dynamic_array_cells(TEMPLATE)}")
    print(f"template parts with metadata: {[p for p in zip_parts(TEMPLATE) if 'metadata' in p]}")

    pristine = OUT / "pristine.xlsx"
    shutil.copyfile(TEMPLATE, pristine)
    report(pristine, "pristine-cache")
    recalc(pristine)
    report(pristine, "pristine-recalc")

    roundtrip = OUT / "roundtrip.xlsx"
    shutil.copyfile(TEMPLATE, roundtrip)
    wb = load_workbook(roundtrip)
    wb.save(roundtrip)
    wb.close()
    print(f"roundtrip parts with metadata: {[p for p in zip_parts(roundtrip) if 'metadata' in p]}")
    recalc(roundtrip)
    report(roundtrip, "roundtrip-recalc")


if __name__ == "__main__":
    main()
