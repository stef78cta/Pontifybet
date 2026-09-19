"""Dump read-only al unor celule din șablonul Cornere V14 (audit).

Rulare: `python scripts/corners_cell_dump.py Dashboard!A26 Verificari!B6 ...`
Fără argumente afișează setul implicit folosit la auditul evaluatorului.
"""

from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from openpyxl import load_workbook  # noqa: E402

from config.settings import TEMPLATES_DIR  # noqa: E402

TEMPLATE = TEMPLATES_DIR / "6_model_analiza_cornere_multiline_optimizat_V14.xlsx"

DEFAULT = [
    *[f"Dashboard!{c}4" for c in "ABCDEFGH"],
    *[f"Dashboard!{c}10" for c in "ABCDEFGH"],
    *[f"Dashboard!{c}25" for c in "ABCDEFGHIKLMNOP"],
    *[f"Dashboard!{c}26" for c in "ABCDEFGHIKLMNOP"],
    *[f"Verificari!{c}{r}" for r in range(5, 16) for c in "ABCD"],
    *[f"Config_v6!{c}{r}" for r in range(1, 62) for c in "AB"],
    *[f"Surse_Import!{c}5" for c in "ABCDEFGHIJKLM"],
    *[f"Surse_Import!{c}6" for c in "ABCDEFGHIJKLM"],
    *[f"Core_Nativ!{c}5" for c in "ABCDEFGHIJKLMNOPQRSTU"],
    *[f"Core_Nativ!{c}{r}" for r in range(6, 20) for c in "ACDHI"],
    *[f"Analiza_Linii!{c}{r}" for r in range(5, 21) for c in "AFGH"],
    *[f"Istoric_Nativ!{c}5" for c in "ABCDEFGHIJKLMNOPQRS"],
    *[f"Calibrare_Linii!{c}{r}" for r in range(5, 20) for c in "ABCDEFGHIJKL"],
    *[f"Input_Meci!{c}5" for c in ("A", "B", "C", "D", "E", "F", "G", "H", "I", "J", "K", "L", "M")],
    *[f"Motor_Meciuri!{c}5" for c in ("A", "B", "C")],
    *[f"Distributie!{c}5" for c in ("A", "B", "C", "D", "E")],
    *[f"OOS_Predictii!{c}5" for c in "ABCDEFGHIJKLMNOPQRSTU"],
]


def main() -> None:
    wanted = sys.argv[1:] or DEFAULT
    wb = load_workbook(TEMPLATE, data_only=False)
    values = load_workbook(TEMPLATE, data_only=True)
    for addr in wanted:
        sheet, coord = addr.split("!")
        raw = wb[sheet][coord].value
        cached = values[sheet][coord].value
        print(f"{addr:26} | {raw!r} | cache={cached!r}")
    wb.close()
    values.close()


if __name__ == "__main__":
    main()
