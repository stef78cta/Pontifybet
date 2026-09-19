"""Audit read-only al formulelor din șablonul Cornere V14.

Inventariază funcțiile, referințele structurate și „forma” fiecărei coloane, ca
extinderea evaluatorului să pornească din formulele reale, nu din presupuneri.
Rulare: `python scripts/corners_formula_audit.py`.
"""

from __future__ import annotations

import re
import sys
from collections import Counter, defaultdict
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from openpyxl import load_workbook  # noqa: E402
from openpyxl.utils.cell import coordinate_from_string  # noqa: E402

from config.settings import TEMPLATES_DIR  # noqa: E402

TEMPLATE = TEMPLATES_DIR / "6_model_analiza_cornere_multiline_optimizat_V14.xlsx"
FUNC_RE = re.compile(r"([A-Za-z_][A-Za-z0-9_.]*)\s*\(")
TABLE_RE = re.compile(r"([A-Za-z_][A-Za-z0-9_]*)\[([^\]]*)\]")


def main() -> None:
    wb = load_workbook(TEMPLATE, data_only=False)
    functions: Counter[str] = Counter()
    tables: Counter[str] = Counter()
    per_column: dict[tuple[str, str], list[tuple[str, str]]] = defaultdict(list)
    total = 0

    for name in wb.sheetnames:
        ws = wb[name]
        for row in ws.iter_rows():
            for cell in row:
                value = cell.value
                if not (isinstance(value, str) and value.startswith("=")):
                    continue
                total += 1
                for fn in FUNC_RE.findall(value):
                    if fn.upper() not in {"IF", "AND", "OR"} or True:
                        functions[fn] += 1
                for tbl, col in TABLE_RE.findall(value):
                    tables[f"{tbl}[{col}]"] += 1
                col_letter, row_idx = coordinate_from_string(cell.coordinate)
                per_column[(name, col_letter)].append((cell.coordinate, value))

    print(f"formule totale: {total}")
    print("\n== funcții ==")
    for fn, count in sorted(functions.items(), key=lambda kv: -kv[1]):
        print(f"{fn:26} {count}")
    print("\n== referințe structurate ==")
    for ref, count in sorted(tables.items()):
        print(f"{ref:44} {count}")
    print("\n== prima formulă per (foaie, coloană) ==")
    for (sheet, col), items in per_column.items():
        coord, formula = items[0]
        print(f"\n--- {sheet}!{col} ({len(items)} formule) @ {coord}")
        print(formula)
    wb.close()


if __name__ == "__main__":
    main()
