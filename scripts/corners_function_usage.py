"""Extrage exemple concrete de utilizare pentru funcțiile V14 neacoperite de evaluator.

Rulează: python scripts/corners_function_usage.py [FUNC ...]
"""

from __future__ import annotations

import re
import sys
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from config.settings import TEMPLATES_DIR  # noqa: E402

WORKBOOK = TEMPLATES_DIR / "6_model_analiza_cornere_multiline_optimizat_V14.xlsx"

TARGETS = sys.argv[1:] or [
    "SUMPRODUCT",
    "_xlfn._xlws.FILTER",
    "LARGE",
    "INT",
    "LEFT",
    "ROWS",
    "COUNTA",
    "TEXT",
]


def main() -> None:
    wb = load_workbook(WORKBOOK, data_only=False, read_only=True)
    samples: dict[str, dict[str, str]] = defaultdict(dict)
    text_formats: Counter[str] = Counter()
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                value = cell.value
                if not isinstance(value, str) or not value.startswith("="):
                    continue
                for func in TARGETS:
                    if func + "(" in value:
                        key = f"{ws.title}!{cell.column_letter}"
                        samples[func].setdefault(key, f"{cell.coordinate}: {value}")
                for match in re.finditer(r'TEXT\([^,]+,"([^"]*)"', value):
                    text_formats[match.group(1)] += 1
    wb.close()

    for func in TARGETS:
        print(f"\n===== {func} =====")
        for key, formula in list(samples.get(func, {}).items())[:6]:
            print(f"--- {key} @ {formula}")

    print("\n===== formate TEXT =====")
    for fmt, count in text_formats.most_common():
        print(f"{fmt!r:28} {count}")


if __name__ == "__main__":
    main()
