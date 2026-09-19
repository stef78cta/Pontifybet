"""Raportează, pentru fiecare foaie V14, coloanele de input față de cele cu formule.

Rulează: python scripts/corners_input_map.py
"""

from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.cell import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from config.settings import TEMPLATES_DIR  # noqa: E402

WORKBOOK = TEMPLATES_DIR / "6_model_analiza_cornere_multiline_optimizat_V14.xlsx"


def main() -> None:
    wb = load_workbook(WORKBOOK, data_only=False)
    for ws in wb.worksheets:
        print(f"\n===== {ws.title} (rânduri={ws.max_row}, coloane={ws.max_column})")
        for name in ws.tables:
            table = ws.tables[name]
            print(f"  tabel {name} ref={table.ref}")
        header_row = 5 if ws.tables else None
        probe_rows = [6, 7] if header_row else list(range(1, min(ws.max_row, 62) + 1))
        for row in probe_rows:
            if row > ws.max_row:
                continue
            formulas: list[str] = []
            literals: list[str] = []
            for col in range(1, ws.max_column + 1):
                value = ws.cell(row=row, column=col).value
                letter = get_column_letter(col)
                if isinstance(value, str) and value.startswith("="):
                    formulas.append(letter)
                elif value is not None:
                    literals.append(f"{letter}={value!r}"[:60])
            print(f"  rând {row}: formule={','.join(formulas)}")
            if literals:
                print(f"           literali={literals}")
    wb.close()


if __name__ == "__main__":
    main()
