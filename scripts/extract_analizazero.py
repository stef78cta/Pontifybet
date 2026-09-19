"""Extrage meciurile 0-0 la pauză din cele 50 de ligi alese, sezoanele 2023/24–2025/26.

Sursa este `league-matches` (același payload ca `/match` pentru statisticile scalare).
FootyStats nu publică xG, șuturi, posesie sau atacuri pe repriză; acele coloane
rămân marcate ca FT-doar, fără valori inventate.
"""

from __future__ import annotations

import json
import sys
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

ROOT = Path(__file__).resolve().parent.parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from src.api.client import FootyStatsClient, FootyStatsError  # noqa: E402

TZ = ZoneInfo("Europe/Bucharest")
SENTINELS = {None, "", -1, -2, "-1", "-2"}
SPLIT_YEARS = {"20232024", "20242025", "20252026"}
CALENDAR_YEARS = {"2024", "2025", "2026"}
OUT_XLSX = ROOT / "analizazero.xlsx"
OUT_XLS = ROOT / "analizazero.xls"
CHECKPOINT = ROOT / "audit" / "analizazero_checkpoint.json"

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
NOTE_FILL = PatternFill("solid", fgColor="FFF2CC")
WRAP = Alignment(wrap_text=True, vertical="center")


def _year_token(raw: Any) -> str:
    return str(raw or "").strip().replace("-", "/").replace(" ", "").replace("/", "")


def season_wanted(year: Any) -> bool:
    token = _year_token(year)
    return token in SPLIT_YEARS or token in CALENDAR_YEARS


def _is_zero(raw: Any) -> bool:
    if raw in SENTINELS or isinstance(raw, bool):
        return False
    try:
        return int(float(raw)) == 0
    except (TypeError, ValueError):
        return False


def is_ht_0_0(match: dict[str, Any]) -> bool:
    """0-0 HT doar cu valori reale; -1/-2 nu sunt zero."""
    return _is_zero(match.get("ht_goals_team_a")) and _is_zero(match.get("ht_goals_team_b"))


def _cell(raw: Any) -> Any:
    if raw is None or raw == "":
        return None
    if isinstance(raw, (int, float, str)) and raw in SENTINELS:
        return None
    if isinstance(raw, bool):
        return raw
    if isinstance(raw, (int, float, str)):
        return raw
    if isinstance(raw, list):
        if not raw:
            return ""
        if all(isinstance(x, (str, int, float)) for x in raw):
            return ", ".join(str(x) for x in raw)
        return json.dumps(raw, ensure_ascii=False, default=str)
    if isinstance(raw, dict):
        return json.dumps(raw, ensure_ascii=False, default=str)
    return str(raw)


def _kickoff(match: dict[str, Any]) -> tuple[str | None, str | None]:
    raw = match.get("date_unix")
    if raw in SENTINELS:
        return None, None
    try:
        dt = datetime.fromtimestamp(int(raw), tz=timezone.utc).astimezone(TZ)
    except (TypeError, ValueError, OSError):
        return None, None
    return dt.strftime("%Y-%m-%d"), dt.strftime("%H:%M")


def _num(raw: Any) -> int | float | None:
    if raw in SENTINELS or isinstance(raw, bool):
        return None
    if isinstance(raw, int):
        return raw
    if isinstance(raw, float):
        return raw
    try:
        text = str(raw).strip().replace(",", ".")
        return int(text) if text.isdigit() else float(text)
    except (TypeError, ValueError):
        return None


def target_seasons(leagues: list[dict[str, Any]]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    seen: set[int] = set()
    for lg in leagues:
        label = str(lg.get("name") or lg.get("league_name") or "")
        country = str(lg.get("country") or "")
        for season in lg.get("season") or []:
            if not isinstance(season, dict) or not season_wanted(season.get("year")):
                continue
            sid = season.get("id")
            if sid in SENTINELS:
                continue
            season_id = int(sid)
            if season_id in seen:
                continue
            seen.add(season_id)
            rows.append(
                {
                    "league": label,
                    "country": country or str(season.get("country") or ""),
                    "season_id": season_id,
                    "season_year": season.get("year"),
                    "season_label": str(season.get("year")),
                }
            )
    rows.sort(key=lambda r: (r["country"], r["league"], str(r["season_year"])))
    return rows


def identity_row(match: dict[str, Any], meta: dict[str, Any]) -> dict[str, Any]:
    date_iso, time_hm = _kickoff(match)
    return {
        "match_id": match.get("id"),
        "data": date_iso,
        "ora_Bucharest": time_hm,
        "liga": meta["league"],
        "tara": meta["country"] or match.get("country") or "",
        "sezon": match.get("season") or meta["season_label"],
        "season_id": meta["season_id"],
        "etapa": match.get("game_week"),
        "status": match.get("status"),
        "gazda": match.get("home_name"),
        "oaspete": match.get("away_name"),
        "home_id": match.get("homeID"),
        "away_id": match.get("awayID"),
        "stadion": match.get("stadium_name") or None,
        "url_meci": match.get("match_url"),
        "ht_gazda": _num(match.get("ht_goals_team_a")),
        "ht_oaspete": _num(match.get("ht_goals_team_b")),
        "ht_total": _num(match.get("HTGoalCount")),
        "ft_gazda": _num(match.get("homeGoalCount")),
        "ft_oaspete": _num(match.get("awayGoalCount")),
        "goluri_2h_gazda": _num(match.get("goals_2hg_team_a")),
        "goluri_2h_oaspete": _num(match.get("goals_2hg_team_b")),
        "goluri_2h_total": _num(match.get("GoalCount_2hg")),
        "minute_goluri_gazda": _cell(match.get("homeGoals")),
        "minute_goluri_oaspete": _cell(match.get("awayGoals")),
    }


NA_XG = "API nu oferă xG pe repriză (doar FT)"
NA_SHOTS = "API nu oferă șuturi pe repriză (doar FT)"
NA_POSS = "API nu oferă posesie pe repriză (doar FT)"
NA_ATT = "API nu oferă atacuri pe repriză (doar FT)"
NA_YR = "Galbene/roșii pe repriză nu sunt separate; doar total cartonașe 1H/2H + galbene/roșii FT"


def half_row(match: dict[str, Any], meta: dict[str, Any], half: int) -> dict[str, Any]:
    row = identity_row(match, meta)
    if half == 1:
        row.update(
            {
                "repriza": 1,
                "goluri_gazda": _num(match.get("ht_goals_team_a")),
                "goluri_oaspete": _num(match.get("ht_goals_team_b")),
                "goluri_total": _num(match.get("HTGoalCount")),
                "cornere_gazda": _num(match.get("team_a_fh_corners")),
                "cornere_oaspete": _num(match.get("team_b_fh_corners")),
                "cornere_total": _num(match.get("corner_fh_count")),
                "cartonase_gazda": _num(match.get("team_a_fh_cards")),
                "cartonase_oaspete": _num(match.get("team_b_fh_cards")),
                "cartonase_total": _num(match.get("total_fh_cards")),
                "cote_1x2_gazda": _num(match.get("odds_1st_half_result_1")),
                "cote_1x2_egal": _num(match.get("odds_1st_half_result_x")),
                "cote_1x2_oaspete": _num(match.get("odds_1st_half_result_2")),
            }
        )
    else:
        row.update(
            {
                "repriza": 2,
                "goluri_gazda": _num(match.get("goals_2hg_team_a")),
                "goluri_oaspete": _num(match.get("goals_2hg_team_b")),
                "goluri_total": _num(match.get("GoalCount_2hg")),
                "cornere_gazda": _num(match.get("team_a_2h_corners")),
                "cornere_oaspete": _num(match.get("team_b_2h_corners")),
                "cornere_total": _num(match.get("corner_2h_count")),
                "cartonase_gazda": _num(match.get("team_a_2h_cards")),
                "cartonase_oaspete": _num(match.get("team_b_2h_cards")),
                "cartonase_total": _num(match.get("total_2h_cards")),
                "cote_1x2_gazda": _num(match.get("odds_2nd_half_result_1")),
                "cote_1x2_egal": _num(match.get("odds_2nd_half_result_x")),
                "cote_1x2_oaspete": _num(match.get("odds_2nd_half_result_2")),
            }
        )
    row.update(
        {
            "xg_gazda_FT": _num(match.get("team_a_xg")),
            "xg_oaspete_FT": _num(match.get("team_b_xg")),
            "xg_total_FT": _num(match.get("total_xg")),
            "suturi_gazda_FT": _num(match.get("team_a_shots")),
            "suturi_poarta_gazda_FT": _num(match.get("team_a_shotsOnTarget")),
            "suturi_langa_gazda_FT": _num(match.get("team_a_shotsOffTarget")),
            "suturi_oaspete_FT": _num(match.get("team_b_shots")),
            "suturi_poarta_oaspete_FT": _num(match.get("team_b_shotsOnTarget")),
            "suturi_langa_oaspete_FT": _num(match.get("team_b_shotsOffTarget")),
            "atacuri_gazda_FT": _num(match.get("team_a_attacks")),
            "atacuri_oaspete_FT": _num(match.get("team_b_attacks")),
            "atacuri_periculoase_gazda_FT": _num(match.get("team_a_dangerous_attacks")),
            "atacuri_periculoase_oaspete_FT": _num(match.get("team_b_dangerous_attacks")),
            "posesie_gazda_FT": _num(match.get("team_a_possession")),
            "posesie_oaspete_FT": _num(match.get("team_b_possession")),
            "galbene_gazda_FT": _num(match.get("team_a_yellow_cards")),
            "galbene_oaspete_FT": _num(match.get("team_b_yellow_cards")),
            "rosii_gazda_FT": _num(match.get("team_a_red_cards")),
            "rosii_oaspete_FT": _num(match.get("team_b_red_cards")),
            "cartonase_num_gazda_FT": _num(match.get("team_a_cards_num")),
            "cartonase_num_oaspete_FT": _num(match.get("team_b_cards_num")),
            "faulturi_gazda_FT": _num(match.get("team_a_fouls")),
            "faulturi_oaspete_FT": _num(match.get("team_b_fouls")),
            "offsides_gazda_FT": _num(match.get("team_a_offsides")),
            "offsides_oaspete_FT": _num(match.get("team_b_offsides")),
            "cornere_gazda_FT": _num(match.get("team_a_corners")),
            "cornere_oaspete_FT": _num(match.get("team_b_corners")),
            "cornere_total_FT": _num(match.get("totalCornerCount")),
            "nota_xg": NA_XG,
            "nota_suturi": NA_SHOTS,
            "nota_posesie": NA_POSS,
            "nota_atacuri": NA_ATT,
            "nota_cartonase_culoare": NA_YR,
        }
    )
    return row


def raw_row(match: dict[str, Any], meta: dict[str, Any]) -> dict[str, Any]:
    row = identity_row(match, meta)
    for key, value in match.items():
        if key in row:
            continue
        row[f"api_{key}"] = _cell(value)
    return row


def _style_header(ws: Worksheet) -> None:
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 32


def _autosize(ws: Worksheet, max_width: int = 28) -> None:
    for idx, col in enumerate(ws.columns, start=1):
        header = str(col[0].value or "")
        width = min(max(12, len(header) + 2), max_width)
        ws.column_dimensions[get_column_letter(idx)].width = width


def write_sheet(wb: Workbook, title: str, rows: list[dict[str, Any]], note_cols: set[str] | None = None) -> None:
    ws = wb.create_sheet(title)
    if not rows:
        ws.append(["(niciun rând)"])
        return
    headers = list(rows[0].keys())
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h) for h in headers])
    _style_header(ws)
    notes = note_cols or set()
    for col_idx, header in enumerate(headers, start=1):
        if header not in notes:
            continue
        for row_idx in range(2, ws.max_row + 1):
            ws.cell(row_idx, col_idx).fill = NOTE_FILL
            ws.cell(row_idx, col_idx).alignment = WRAP
    _autosize(ws)


def write_notes(wb: Workbook, seasons: list[dict[str, Any]], errors: list[str], totals: dict[str, Any]) -> None:
    ws = wb.create_sheet("Limite_API", 0)
    lines = [
        ["Pontifybet — extract 0-0 la pauză"],
        ["Generat (Europe/Bucharest)", datetime.now(TZ).strftime("%Y-%m-%d %H:%M")],
        ["Sursă", "FootyStats /league-matches (ligi alese pe cheia API)"],
        ["Filtru", "status=complete ȘI ht_goals_team_a=0 ȘI ht_goals_team_b=0"],
        ["Sezoane", "2023/2024, 2024/2025, 2025/2026 (calendar: 2024, 2025, 2026)"],
        ["Ligi în listă", totals.get("leagues")],
        ["Sezoane țintă", len(seasons)],
        ["Sezoane încărcate", totals.get("seasons_ok")],
        ["Meciuri complete scanate", totals.get("complete_scanned")],
        ["Meciuri 0-0 HT", totals.get("ht00")],
        [],
        ["Ce există nativ pe repriză"],
        ["Goluri", "ht_goals_* / goals_2hg_*"],
        ["Cornere", "team_*_fh_corners / team_*_2h_corners"],
        ["Cartonașe (total, nu culoare)", "team_*_fh_cards / team_*_2h_cards"],
        ["Cote 1X2 pe repriză", "odds_1st_half_result_* / odds_2nd_half_result_*"],
        [],
        ["Ce NU există pe repriză (doar full-time)"],
        ["xG", NA_XG],
        ["Șuturi (total / pe poartă / pe lângă)", NA_SHOTS],
        ["Posesie", NA_POSS],
        ["Atacuri / atacuri periculoase", NA_ATT],
        ["Galbene vs roșii pe repriză", NA_YR],
        [],
        ["Foi"],
        ["Repriza_1", "toate meciurile 0-0 HT, metrici 1H native + FT-doar etichetate"],
        ["Repriza_2", "aceleași meciuri, metrici 2H native + FT-doar etichetate"],
        ["Toate_campurile_API", "payload scalar complet league-matches, fără câmpuri inventate"],
        ["Rezumat_ligi", "număr 0-0 HT pe ligă/sezon"],
        [],
        ["Erori pe sezoane"],
    ]
    for line in lines:
        ws.append(line)
    if errors:
        for err in errors:
            ws.append([err])
    else:
        ws.append(["(niciuna)"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.column_dimensions["A"].width = 36
    ws.column_dimensions["B"].width = 88


def write_summary(wb: Workbook, summary: list[dict[str, Any]]) -> None:
    write_sheet(wb, "Rezumat_ligi", summary or [{"liga": None, "sezon": None, "ht00": 0}])


def save_workbook(wb: Workbook) -> None:
    """xlsx este formatul real; .xls e copie cu numele cerut (Excel 2007+)."""
    OUT_XLSX.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT_XLSX)
    wb.save(OUT_XLS)


def fetch_all() -> None:
    sys.stdout.reconfigure(encoding="utf-8")
    CHECKPOINT.parent.mkdir(parents=True, exist_ok=True)
    errors: list[str] = []
    half1: list[dict[str, Any]] = []
    half2: list[dict[str, Any]] = []
    raw_rows: list[dict[str, Any]] = []
    summary: list[dict[str, Any]] = []
    complete_scanned = 0
    seasons_ok = 0

    with FootyStatsClient() as client:
        leagues = client.list_leagues(chosen_only=True)
        seasons = target_seasons(leagues)
        print(f"ligi={len(leagues)} sezoane_tinta={len(seasons)}")
        for idx, meta in enumerate(seasons, start=1):
            sid = meta["season_id"]
            label = f"{meta['league']} {meta['season_label']} ({sid})"
            try:
                matches = client.league_matches(sid, all_pages=True)
            except FootyStatsError as exc:
                msg = f"FAIL {label}: {exc.user_message}"
                errors.append(msg)
                print(msg)
                summary.append(
                    {
                        "liga": meta["league"],
                        "tara": meta["country"],
                        "sezon": meta["season_label"],
                        "season_id": sid,
                        "meciuri_complete": None,
                        "ht00": None,
                        "eroare": exc.user_message,
                    }
                )
                continue
            seasons_ok += 1
            ht00_n = 0
            complete_n = 0
            for match in matches:
                if not isinstance(match, dict):
                    continue
                if str(match.get("status") or "").lower() != "complete":
                    continue
                complete_n += 1
                complete_scanned += 1
                if not is_ht_0_0(match):
                    continue
                ht00_n += 1
                half1.append(half_row(match, meta, 1))
                half2.append(half_row(match, meta, 2))
                raw_rows.append(raw_row(match, meta))
            summary.append(
                {
                    "liga": meta["league"],
                    "tara": meta["country"],
                    "sezon": meta["season_label"],
                    "season_id": sid,
                    "meciuri_complete": complete_n,
                    "ht00": ht00_n,
                    "eroare": None,
                }
            )
            print(f"[{idx}/{len(seasons)}] {label}: complete={complete_n} ht00={ht00_n}")

    totals = {
        "leagues": len(leagues),
        "seasons_ok": seasons_ok,
        "complete_scanned": complete_scanned,
        "ht00": len(half1),
    }
    CHECKPOINT.write_text(
        json.dumps({"totals": totals, "errors": errors, "summary": summary}, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )

    half1.sort(key=lambda r: (str(r.get("data") or ""), str(r.get("liga") or ""), str(r.get("gazda") or "")))
    order = {(r["match_id"], r["data"], r["liga"]) : i for i, r in enumerate(half1)}
    half2.sort(key=lambda r: order.get((r["match_id"], r["data"], r["liga"]), 10**9))
    raw_rows.sort(key=lambda r: order.get((r["match_id"], r["data"], r["liga"]), 10**9))

    wb = Workbook()
    default = wb.active
    wb.remove(default)
    write_notes(wb, seasons, errors, totals)
    note_cols = {"nota_xg", "nota_suturi", "nota_posesie", "nota_atacuri", "nota_cartonase_culoare"}
    write_sheet(wb, "Repriza_1", half1, note_cols)
    write_sheet(wb, "Repriza_2", half2, note_cols)
    write_sheet(wb, "Toate_campurile_API", raw_rows)
    write_summary(wb, summary)
    save_workbook(wb)
    print(f"SCRIS {OUT_XLS} | {OUT_XLSX} | meciuri_0-0={len(half1)}")


if __name__ == "__main__":
    fetch_all()
