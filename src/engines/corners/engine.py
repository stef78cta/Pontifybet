"""API public al motorului Cornere Multi-Line V14.

Motorul nu reimplementează metodologia: încarcă șablonul V14 read-only, aplică
override-uri pe celulele de input și cere `WorkbookEngine` valorile produse de
formulele originale. Un singur calcul per lot alimentează cele 14 linii ale
fiecărui meci, plus `Motor_Meciuri`, `Core_Nativ`, `Distributie` și `Dashboard`
pentru audit.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from functools import lru_cache
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook
from openpyxl.utils.cell import get_column_letter, range_boundaries
from openpyxl.workbook.workbook import Workbook

from config.settings import TEMPLATES_DIR
from src.engines.corners.inputs import (
    CornersBatchInput,
    CornersCapacityError,
    build_overrides,
)
from src.engines.corners.lines import CORNERS_LINES, LINES_PER_SLOT, CornersLine
from src.engines.over05.evaluator import (
    FormulaSupportError,
    WorkbookEngine,
    XLException,
)

TEMPLATE_NAME = "6_model_analiza_cornere_multiline_optimizat_V14.xlsx"

LINES_SHEET = "Analiza_Linii"
MOTOR_SHEET = "Motor_Meciuri"
CORE_SHEET = "Core_Nativ"
DISTRIBUTION_SHEET = "Distributie"
DASHBOARD_SHEET = "Dashboard"

#: Primul rând de date al fiecărui tabel de rezultate și pasul pe slot.
LINES_FIRST_ROW = 6
MOTOR_FIRST_ROW = 6
CORE_FIRST_ROW = 6
CORE_ROWS_PER_SLOT = 12
DISTRIBUTION_FIRST_ROW = 6
DASHBOARD_FIRST_ROW = 26

#: `Dashboard` afișează maximum zece selecții LIVE (rândurile 11–20).
TOP_LIVE_LIMIT = 10

#: Zona liberă de pe `Input_Meci`, sub tabelul `MatchesTable` (care se termină la
#: rândul 105). Nu conține formule, deci exportul poate scrie rezultatele Python
#: fără să atingă niciun calcul al workbook-ului.
RESULT_ZONE_SHEET = "Input_Meci"
RESULT_ZONE_FIRST_ROW = 107
RESULT_ZONE_COLUMNS = 17  # A–Q


class CornersEngineError(Exception):
    """Eroare tehnică de evaluare, distinctă de stările metodologice ale modelului.

    Un `NOT RELEASED`, `G0 FAIL`, `WATCH` sau `NO BET` este un rezultat valid al
    modelului. Această excepție semnalează că evaluatorul nu a putut produce un
    rezultat (formulă nesuportată, referință invalidă, ciclu).
    """

    def __init__(self, message: str, *, cell: str | None = None) -> None:
        super().__init__(message)
        self.message = message
        self.cell = cell


@lru_cache(maxsize=1)
def load_corners_workbook(path: str | None = None) -> Workbook:
    """Încarcă o singură dată șablonul V14; workbook-ul este folosit read-only."""
    workbook_path = Path(path) if path else TEMPLATES_DIR / TEMPLATE_NAME
    return load_workbook(workbook_path, data_only=False)


@lru_cache(maxsize=8)
def _table_columns(sheet: str, path: str | None = None) -> dict[str, str]:
    """Mapează numele coloanelor unui tabel Excel la litere, direct din șablon."""
    workbook = load_corners_workbook(path)
    worksheet = workbook[sheet]
    name = next(iter(worksheet.tables))
    table = worksheet.tables[name]
    min_col = range_boundaries(table.ref)[0]
    return {
        column.name: get_column_letter(min_col + offset)
        for offset, column in enumerate(table.tableColumns)
    }


def lines_columns() -> dict[str, str]:
    return _table_columns(LINES_SHEET)


def motor_columns() -> dict[str, str]:
    return _table_columns(MOTOR_SHEET)


def core_columns() -> dict[str, str]:
    return _table_columns(CORE_SHEET)


def distribution_columns() -> dict[str, str]:
    return _table_columns(DISTRIBUTION_SHEET)


def _as_number_or_none(value: Any) -> float | None:
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    return None


def _as_int_if_whole(value: Any) -> int | float | None:
    number = _as_number_or_none(value)
    if number is None:
        return None
    return int(number) if number == int(number) else number


def _as_text(value: Any) -> str | None:
    if value is None or value == "":
        return None
    return str(value)


@dataclass
class CornersLineResult:
    """Rezultatul oficial al unei linii, cu toate coloanele din `LinesTable`."""

    line: CornersLine
    excel_row: int
    cells: dict[str, Any]
    """Valorile pe nume de coloană din tabel, exact cum le-a produs workbook-ul."""

    rank_live: int | None = None
    """Rangul Top LIVE global; vine din `Rank_LIVE` sau din selectorul echivalent."""

    rank_live_source: str = "workbook"

    def value(self, column: str) -> Any:
        return self.cells.get(column)

    @property
    def label(self) -> str:
        return self.line.label

    @property
    def release(self) -> str | None:
        return _as_text(self.cells.get("Release"))

    @property
    def g0(self) -> str | None:
        return _as_text(self.cells.get("G0"))

    @property
    def mu(self) -> float | None:
        return _as_number_or_none(self.cells.get("Mu"))

    @property
    def dispersion(self) -> float | None:
        return _as_number_or_none(self.cells.get("D"))

    @property
    def p_raw(self) -> float | None:
        return _as_number_or_none(self.cells.get("P_Raw"))

    @property
    def p_calibrated(self) -> float | None:
        return _as_number_or_none(self.cells.get("P_Calibrated"))

    @property
    def p_final(self) -> float | None:
        return _as_number_or_none(self.cells.get("P_FINAL"))

    @property
    def failure_gate(self) -> float | None:
        return _as_number_or_none(self.cells.get("Failure_Gate"))

    @property
    def failure_model(self) -> float | None:
        return _as_number_or_none(self.cells.get("Failure_Model"))

    @property
    def risk_score(self) -> float | None:
        return _as_number_or_none(self.cells.get("Risk_Score"))

    @property
    def risk_level_raw(self) -> int | float | None:
        return _as_int_if_whole(self.cells.get("Risk_Level_Raw"))

    @property
    def risk_level(self) -> int | float | None:
        return _as_int_if_whole(self.cells.get("Risk_Level_FINAL"))

    @property
    def confidence_final(self) -> float | None:
        return _as_number_or_none(self.cells.get("Confidence_Final"))

    @property
    def candidate_confidence(self) -> float | None:
        return _as_number_or_none(self.cells.get("Candidate_Confidence"))

    @property
    def tail_gate(self) -> str | None:
        return _as_text(self.cells.get("Tail_Gate"))

    @property
    def defensive_candidate(self) -> str | None:
        return _as_text(self.cells.get("Defensive_Candidate"))

    @property
    def defensive_gate(self) -> str | None:
        return _as_text(self.cells.get("Defensive_Gate"))

    @property
    def verdict(self) -> str | None:
        return _as_text(self.cells.get("Verdict_FINAL"))

    @property
    def rank_in_match(self) -> int | None:
        value = _as_int_if_whole(self.cells.get("Rank_In_Match"))
        return int(value) if isinstance(value, int) else None

    @property
    def eligible(self) -> bool:
        return _as_number_or_none(self.cells.get("Eligible")) == 1.0

    @property
    def builder(self) -> str | None:
        return _as_text(self.cells.get("Builder"))

    @property
    def live_eligible_best(self) -> bool:
        return _as_number_or_none(self.cells.get("Live_Eligible_Best")) == 1.0

    @property
    def reason(self) -> str | None:
        return _as_text(self.cells.get("Reason"))

    @property
    def model_signature(self) -> str | None:
        return _as_text(self.cells.get("Model_Signature"))

    @property
    def oos_status(self) -> str | None:
        return _as_text(self.cells.get("OOS_Status"))

    def payload(self) -> dict[str, Any]:
        """Formă serializabilă, folosită de UI, istoric și CSV."""
        return {
            "line": self.line.label,
            "market_type": self.line.market_type,
            "k": self.line.k,
            "line_value": self.line.line,
            "excel_row": self.excel_row,
            "release": self.release,
            "g0": self.g0,
            "mu": self.mu,
            "d": self.dispersion,
            "p_raw": self.p_raw,
            "p_calibrated": self.p_calibrated,
            "p_final": self.p_final,
            "failure_model": self.failure_model,
            "failure_gate": self.failure_gate,
            "risk_score": self.risk_score,
            "risk_level_raw": self.risk_level_raw,
            "risk_level": self.risk_level,
            "confidence_final": self.confidence_final,
            "candidate_confidence": self.candidate_confidence,
            "tail_gate": self.tail_gate,
            "defensive_candidate": self.defensive_candidate,
            "defensive_gate": self.defensive_gate,
            "verdict": self.verdict,
            "rank_in_match": self.rank_in_match,
            "eligible": self.eligible,
            "builder": self.builder,
            "live_eligible_best": self.live_eligible_best,
            "rank_live": self.rank_live,
            "rank_live_source": self.rank_live_source,
            "reason": self.reason,
            "oos_status": self.oos_status,
            "model_signature": self.model_signature,
            "empirical_ha": _as_number_or_none(self.cells.get("Empirical_HA")),
            "empirical_l10": _as_number_or_none(self.cells.get("Empirical_L10")),
            "empirical_l20": _as_number_or_none(self.cells.get("Empirical_L20")),
            "n_ha_empirical": _as_number_or_none(self.cells.get("N_HA_Empirical")),
            "n_l10_empirical": _as_number_or_none(self.cells.get("N_L10_Empirical")),
            "n_l20_empirical": _as_number_or_none(self.cells.get("N_L20_Empirical")),
        }


@dataclass
class CornersMatchResult:
    """Rezultatul oficial al unui meci: cele 14 linii plus intermediarii de audit."""

    match_id: str
    slot: int
    lines: tuple[CornersLineResult, ...]
    motor: dict[str, Any] = field(default_factory=dict)
    core: tuple[dict[str, Any], ...] = ()
    distribution: dict[str, Any] = field(default_factory=dict)
    dashboard: dict[str, Any] = field(default_factory=dict)

    @property
    def mode(self) -> str | None:
        return _as_text(self.motor.get("Mode"))

    @property
    def release(self) -> str | None:
        return _as_text(self.motor.get("Release"))

    @property
    def g0(self) -> str | None:
        return _as_text(self.motor.get("G0"))

    @property
    def reason(self) -> str | None:
        return _as_text(self.motor.get("Reason"))

    @property
    def match_label(self) -> str | None:
        return _as_text(self.motor.get("Match"))

    @property
    def distribution_kind(self) -> str | None:
        return _as_text(self.motor.get("Distribution"))

    @property
    def selected_line(self) -> CornersLineResult | None:
        """Linia cu `Rank_In_Match = 1`, adică selecția pe care o arată Dashboard."""
        for item in self.lines:
            if item.rank_in_match == 1:
                return item
        return None

    def by_label(self, label: str) -> CornersLineResult:
        for item in self.lines:
            if item.line.label == label:
                return item
        raise KeyError(label)

    def summary_recommendation(self) -> str:
        """Verdictul afișat pentru meci, în logica `Dashboard!G`."""
        selected = self.selected_line
        if self.release != "READY":
            return "NOT RELEASED"
        if selected is None:
            return "NO BET" if self.g0 == "FAIL" else "WATCH / NO BET"
        return selected.verdict or "WATCH / NO BET"

    def payload(self) -> list[dict[str, Any]]:
        return [item.payload() for item in self.lines]


@dataclass
class CornersBatchResult:
    """Rezultatele unui lot evaluat într-o singură trecere prin workbook."""

    matches: dict[str, CornersMatchResult]
    batch: CornersBatchInput
    extras: dict[str, Any] = field(default_factory=dict)
    """Celule suplimentare cerute prin `extra_cells`, pentru audit și diagnostic."""

    def ordered(self) -> list[CornersMatchResult]:
        return [self.matches[m.match_id] for m in self.batch.matches if m.match_id in self.matches]


def _parse_addr(addr: str) -> tuple[str, str]:
    sheet, coord = addr.split("!", 1)
    return sheet.strip("'").replace("''", "'"), coord.replace("$", "")


def _read(engine: WorkbookEngine, sheet: str, coord: str) -> Any:
    """Citește o celulă, separând erorile de model de defectele evaluatorului."""
    try:
        return engine.cell(sheet, coord)
    except FormulaSupportError as exc:
        raise CornersEngineError(
            f"Construcție Excel nesuportată în {sheet}!{coord}: {exc}",
            cell=f"{sheet}!{coord}",
        ) from exc
    except RecursionError as exc:
        raise CornersEngineError(
            f"Adâncime de evaluare depășită în {sheet}!{coord}",
            cell=f"{sheet}!{coord}",
        ) from exc
    except (XLException, ZeroDivisionError, ValueError, OverflowError) as exc:
        # Erorile Excel propagate până la o celulă de ieșire sunt un defect de
        # integrare: formulele V14 își tratează singure cazurile de model.
        raise CornersEngineError(
            f"Eroare de evaluare în {sheet}!{coord}: {type(exc).__name__}",
            cell=f"{sheet}!{coord}",
        ) from exc


def _read_row(
    engine: WorkbookEngine,
    sheet: str,
    row: int,
    columns: dict[str, str],
) -> dict[str, Any]:
    return {name: _read(engine, sheet, f"{letter}{row}") for name, letter in columns.items()}


def compute_corners_batch(
    batch: CornersBatchInput,
    *,
    workbook_path: str | Path | None = None,
    extra_cells: Iterable[str] | None = None,
) -> CornersBatchResult:
    """Evaluează un lot de meciuri pe șablonul V14.

    @param batch - meciurile, istoricul nativ și sursele care le acoperă.
    @param workbook_path - șablon alternativ (folosit de teste de paritate).
    @param extra_cells - celule suplimentare de evaluat, în notație `Foaie!A1`.
    @returns - rezultatele celor 14 linii per meci plus intermediarii de audit.
    @raises CornersEngineError - defect tehnic de evaluare.
    @raises CornersCapacityError - lotul depășește capacitatea șablonului.
    """
    if not batch.matches:
        return CornersBatchResult(matches={}, batch=batch)

    workbook = load_corners_workbook(str(workbook_path) if workbook_path else None)
    overrides = build_overrides(batch)
    engine = WorkbookEngine(workbook, overrides, dates_as_serial=True)

    line_cols = lines_columns()
    motor_cols = motor_columns()
    core_cols = core_columns()
    dist_cols = distribution_columns()

    results: dict[str, CornersMatchResult] = {}
    for index, match in enumerate(batch.matches):
        slot = index + 1
        motor_row = MOTOR_FIRST_ROW + index
        dashboard_row = DASHBOARD_FIRST_ROW + index
        first_line_row = LINES_FIRST_ROW + index * LINES_PER_SLOT
        first_core_row = CORE_FIRST_ROW + index * CORE_ROWS_PER_SLOT

        lines: list[CornersLineResult] = []
        for line in CORNERS_LINES:
            excel_row = first_line_row + line.offset
            cells = _read_row(engine, LINES_SHEET, excel_row, line_cols)
            rank = _as_int_if_whole(cells.get("Rank_LIVE"))
            lines.append(
                CornersLineResult(
                    line=line,
                    excel_row=excel_row,
                    cells=cells,
                    rank_live=int(rank) if isinstance(rank, int) else None,
                )
            )

        results[match.match_id] = CornersMatchResult(
            match_id=match.match_id,
            slot=slot,
            lines=tuple(lines),
            motor=_read_row(engine, MOTOR_SHEET, motor_row, motor_cols),
            core=tuple(
                _read_row(engine, CORE_SHEET, first_core_row + offset, core_cols)
                for offset in range(CORE_ROWS_PER_SLOT)
            ),
            distribution=_read_row(
                engine, DISTRIBUTION_SHEET, DISTRIBUTION_FIRST_ROW + index, dist_cols
            ),
            dashboard={
                letter: _read(engine, DASHBOARD_SHEET, f"{letter}{dashboard_row}")
                for letter in ("A", "C", "D", "E", "F", "G", "H", "I", "K", "L", "M", "N", "O", "P")
            },
        )

    extras: dict[str, Any] = {}
    for addr in extra_cells or ():
        sheet, coord = _parse_addr(addr)
        extras[addr] = _read(engine, sheet, coord)

    return CornersBatchResult(matches=results, batch=batch, extras=extras)


def rank_top_live(results: Iterable[CornersMatchResult]) -> list[CornersLineResult]:
    """Ordonează selecțiile LIVE după criteriile din `Dashboard!P`.

    Reproduce exact secvența de departajare a workbook-ului: Risk Level crescător,
    P_FINAL descrescător, Confidence descrescător, apoi ordinea slotului. Este
    necesar doar când analiza s-a împărțit în mai multe loturi, fiecare cu propriul
    `Istoric_Nativ`: în interiorul unui lot, `Rank_LIVE` este deja calculat de
    workbook, iar testele de paritate verifică echivalența celor două căi.

    @returns - selecțiile eligibile, în ordinea clasamentului, câte una pe meci.
    """
    candidates: list[tuple[tuple[Any, ...], CornersLineResult]] = []
    for order, result in enumerate(results):
        if result.mode != "LIVE":
            continue
        selected = result.selected_line
        if selected is None or not selected.eligible:
            continue
        level = _as_number_or_none(selected.cells.get("Risk_Level_FINAL"))
        p_final = _as_number_or_none(selected.cells.get("P_FINAL"))
        confidence = _as_number_or_none(selected.cells.get("Confidence_Final"))
        if level is None or p_final is None or confidence is None:
            continue
        candidates.append(((level, -p_final, -confidence, order), selected))

    candidates.sort(key=lambda item: item[0])
    ordered = [item[1] for item in candidates]
    for position, item in enumerate(ordered, start=1):
        item.rank_live = position
        item.rank_live_source = "python_cross_batch"
    return ordered


def top_live_selection(results: Iterable[CornersMatchResult]) -> list[CornersLineResult]:
    """Primele zece selecții LIVE, ca în `Dashboard` rândurile 11–20."""
    return rank_top_live(results)[:TOP_LIVE_LIMIT]


#: Etichetele blocului de rezultate Python scris în zona liberă de pe `Input_Meci`.
RESULT_ZONE_HEADERS: tuple[str, ...] = (
    "Match_ID",
    "Meci",
    "Mod",
    "Linie",
    "Release",
    "G0",
    "P_Raw",
    "P_FINAL",
    "Failure_Gate",
    "Risk_Score",
    "Confidence_Final",
    "Tail_Gate",
    "Defensive_Gate",
    "Risk_Level_FINAL",
    "Verdict_FINAL",
    "Rank_In_Match",
    "Rank_LIVE",
)


def python_result_writes(
    results: Iterable[CornersMatchResult],
) -> list[dict[str, Any]]:
    """Rezultatele de scris în zona liberă de rezultate, un dict per linie.

    Exportul rămâne opțional: blocul face rezultatele Python vizibile în fișier
    fără recalculare externă, iar toate formulele originale rămân intacte.
    Plasarea concretă (foaie, primul rând, coloane) vine din registry, nu de aici.
    """
    rows: list[dict[str, Any]] = []
    for result in results:
        for item in result.lines:
            rows.append(
                {
                    "Match_ID": result.match_id,
                    "Meci": result.match_label,
                    "Mod": result.mode,
                    "Linie": item.line.label,
                    "Release": item.release,
                    "G0": item.g0,
                    "P_Raw": item.p_raw,
                    "P_FINAL": item.p_final,
                    "Failure_Gate": item.failure_gate,
                    "Risk_Score": item.risk_score,
                    "Confidence_Final": item.confidence_final,
                    "Tail_Gate": item.tail_gate,
                    "Defensive_Gate": item.defensive_gate,
                    "Risk_Level_FINAL": item.risk_level,
                    "Verdict_FINAL": item.verdict,
                    "Rank_In_Match": item.rank_in_match,
                    "Rank_LIVE": item.rank_live,
                }
            )
    return rows


__all__ = [
    "CornersBatchResult",
    "CornersCapacityError",
    "CornersEngineError",
    "CornersLineResult",
    "CornersMatchResult",
    "RESULT_ZONE_COLUMNS",
    "RESULT_ZONE_FIRST_ROW",
    "RESULT_ZONE_HEADERS",
    "RESULT_ZONE_SHEET",
    "TEMPLATE_NAME",
    "compute_corners_batch",
    "load_corners_workbook",
    "python_result_writes",
    "rank_top_live",
    "top_live_selection",
]
