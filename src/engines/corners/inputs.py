"""Contractul de input al motorului Cornere V14 și traducerea lui în override-uri de celule.

V14 nu consumă medii agregate: `Core_Nativ` și `Motor_Meciuri` recalculează totul
din evenimentele native din `Istoric_Nativ`, validate prin `Surse_Import`. Modulul
descrie exact acele intrări și le proiectează pe intervalele de input ale
șablonului, fără să atingă nicio celulă cu formulă.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Any

from openpyxl.utils.cell import get_column_letter

from src.engines.over05.evaluator import to_serial

#: Limitele fizice ale șablonului V14 (rândurile de date ale fiecărui tabel).
MAX_MATCHES = 100
MAX_HISTORY_ROWS = 400
MAX_SOURCES = 8
MAX_OOS_ROWS = 250

INPUT_SHEET = "Input_Meci"
HISTORY_SHEET = "Istoric_Nativ"
SOURCE_SHEET = "Surse_Import"
OOS_SHEET = "OOS_Predictii"

DATA_START_ROW = 6

#: Coloanele de input ale fiecărei foi (restul sunt formule și rămân neatinse).
INPUT_COLUMNS = 33  # Input_Meci A–AG
HISTORY_COLUMNS = 13  # Istoric_Nativ A–M
SOURCE_COLUMNS = 12  # Surse_Import A–L
OOS_COLUMNS = 19  # OOS_Predictii A–S

#: Coloanele de context opțional din `Input_Meci` (soft + contextual), în ordinea
#: șablonului. Sunt inputuri pre-match: absența lor este o stare tratată de model
#: (`Soft_Missing`), nu o eroare.
CONTEXT_FIELDS: tuple[str, ...] = (
    "Shots_H",
    "Shots_A",
    "Blocked_H",
    "Blocked_A",
    "Crosses_H",
    "Crosses_A",
    "Possession_H",
    "Possession_A",
    "Field_Tilt_H",
    "Field_Tilt_A",
    "Attack_H",
    "Attack_A",
    "Chasing_H",
    "Chasing_A",
    "Wings_H",
    "Wings_A",
    "Absences_H",
    "Absences_A",
    "Stakes_H",
    "Stakes_A",
)


class CornersCapacityError(Exception):
    """Inputul depășește capacitatea fizică a șablonului V14.

    Ridicată în loc să trunchieze datele: V14 calculează ferestre și eșantioane
    pe istoricul complet, deci o tăiere silențioasă ar schimba rezultatele.
    """

    def __init__(self, message: str, *, limit: int, requested: int) -> None:
        super().__init__(message)
        self.message = message
        self.limit = limit
        self.requested = requested


@dataclass(frozen=True)
class NativeHistoryRow:
    """Un eveniment istoric nativ: un meci încheiat cu cornerele lui oficiale."""

    event_id: str
    league: str
    season: str
    match_date: date
    home: str
    away: str
    home_corners: int
    away_corners: int
    source_id: str
    source_url: str
    retrieved_utc: datetime
    definition: str
    evidence: str

    @property
    def dedup_key(self) -> tuple[str, str, str, str, str]:
        """Cheia canonică folosită și de `Istoric_Nativ!N` pentru duplicate."""
        return (
            self.league,
            self.season,
            self.match_date.strftime("%Y%m%d"),
            self.home,
            self.away,
        )

    def teams(self) -> tuple[str, str]:
        return self.home, self.away


@dataclass(frozen=True)
class NativeSource:
    """Metadatele unei surse de istoric nativ, așa cum le verifică `Core_Nativ!R`."""

    source_id: str
    league: str
    season: str
    provider: str
    url: str
    retrieved_utc: datetime
    coverage: str
    status: str
    rows: int
    missing_corners: int
    attempts: str
    sha256: str

    @property
    def source_key(self) -> str:
        return f"{self.league}|{self.season}"


@dataclass(frozen=True)
class CornersMatchInput:
    """Un rând din `Input_Meci`, cu identitatea meciului și decizia de sursă."""

    match_id: str
    league: str
    season: str
    home: str
    away: str
    kickoff_utc: datetime
    cutoff_utc: datetime
    mode: str = "LIVE"
    sourcing_decision: str = "AUTO"
    evidence_or_attempts: str = ""
    prior_season: str = ""
    prior_allowed: str = "NO"
    notes: str = ""
    context: dict[str, float | None] = field(default_factory=dict)

    def context_value(self, name: str) -> Any:
        value = self.context.get(name)
        return None if value is None else float(value)


@dataclass(frozen=True)
class CornersBatchInput:
    """Un lot evaluat într-o singură trecere prin workbook.

    Toate meciurile lotului partajează același `Istoric_Nativ`, deci lotul este
    unitatea în care `Dashboard` produce clasamentul Top LIVE nativ.
    """

    matches: tuple[CornersMatchInput, ...]
    history: tuple[NativeHistoryRow, ...]
    sources: tuple[NativeSource, ...]
    oos_records: tuple[tuple[Any, ...], ...] = ()

    def validate_capacity(self) -> None:
        if len(self.matches) > MAX_MATCHES:
            raise CornersCapacityError(
                f"Șablonul V14 acceptă {MAX_MATCHES} meciuri per evaluare, "
                f"lotul cere {len(self.matches)}.",
                limit=MAX_MATCHES,
                requested=len(self.matches),
            )
        if len(self.history) > MAX_HISTORY_ROWS:
            raise CornersCapacityError(
                f"Șablonul V14 acceptă {MAX_HISTORY_ROWS} rânduri de istoric nativ, "
                f"lotul cere {len(self.history)}.",
                limit=MAX_HISTORY_ROWS,
                requested=len(self.history),
            )
        if len(self.sources) > MAX_SOURCES:
            raise CornersCapacityError(
                f"Șablonul V14 acceptă {MAX_SOURCES} surse de import, "
                f"lotul cere {len(self.sources)}.",
                limit=MAX_SOURCES,
                requested=len(self.sources),
            )
        if len(self.oos_records) > MAX_OOS_ROWS:
            raise CornersCapacityError(
                f"Șablonul V14 acceptă {MAX_OOS_ROWS} înregistrări OOS, "
                f"lotul cere {len(self.oos_records)}.",
                limit=MAX_OOS_ROWS,
                requested=len(self.oos_records),
            )

    def slot_of(self, match_id: str) -> int:
        """Slotul (1-based) ocupat de un meci în lot."""
        for index, match in enumerate(self.matches, start=1):
            if match.match_id == match_id:
                return index
        raise KeyError(match_id)


def _blank_block(sheet: str, columns: int, rows: int) -> dict[tuple[str, str], Any]:
    """Golește un interval de input, ca datele demonstrative să nu contamineze rularea."""
    letters = [get_column_letter(index) for index in range(1, columns + 1)]
    return {
        (sheet, f"{letter}{row}"): None
        for row in range(DATA_START_ROW, DATA_START_ROW + rows)
        for letter in letters
    }


def _text_or_blank(value: Any) -> Any:
    """Excel distinge celula goală de șirul gol; inputurile absente rămân goale."""
    if value is None:
        return None
    text = str(value)
    return text if text != "" else None


def build_overrides(batch: CornersBatchInput) -> dict[tuple[str, str], Any]:
    """Proiectează un lot pe celulele de input ale șablonului V14.

    Toate rândurile de input sunt întâi golite: șablonul livrat conține trei
    exemple REPLAY cu istoricul și sursa lor, care altfel ar rămâne în calcul.

    @returns - dicționar `(foaie, coordonată) → valoare`, consumabil direct de
        `WorkbookEngine` și de adaptorul de export.
    """
    batch.validate_capacity()

    overrides: dict[tuple[str, str], Any] = {}
    overrides.update(_blank_block(INPUT_SHEET, INPUT_COLUMNS, MAX_MATCHES))
    overrides.update(_blank_block(HISTORY_SHEET, HISTORY_COLUMNS, MAX_HISTORY_ROWS))
    overrides.update(_blank_block(SOURCE_SHEET, SOURCE_COLUMNS, MAX_SOURCES))
    overrides.update(_blank_block(OOS_SHEET, OOS_COLUMNS, MAX_OOS_ROWS))

    for index, match in enumerate(batch.matches):
        row = DATA_START_ROW + index
        values: dict[str, Any] = {
            "A": match.match_id,
            "B": match.league,
            "C": match.season,
            "D": match.home,
            "E": match.away,
            "F": to_serial(match.kickoff_utc),
            "G": to_serial(match.cutoff_utc),
            "H": match.mode,
            "I": _text_or_blank(match.sourcing_decision),
            "J": _text_or_blank(match.evidence_or_attempts),
            "K": _text_or_blank(match.prior_season),
            "L": _text_or_blank(match.prior_allowed),
            "M": _text_or_blank(match.notes),
        }
        for offset, name in enumerate(CONTEXT_FIELDS):
            values[get_column_letter(14 + offset)] = match.context_value(name)
        for letter, value in values.items():
            overrides[(INPUT_SHEET, f"{letter}{row}")] = value

    for index, event in enumerate(batch.history):
        row = DATA_START_ROW + index
        values = {
            "A": event.event_id,
            "B": event.league,
            "C": event.season,
            "D": to_serial(event.match_date),
            "E": event.home,
            "F": event.away,
            "G": int(event.home_corners),
            "H": int(event.away_corners),
            "I": event.source_id,
            "J": event.source_url,
            "K": to_serial(event.retrieved_utc),
            "L": event.definition,
            "M": event.evidence,
        }
        for letter, value in values.items():
            overrides[(HISTORY_SHEET, f"{letter}{row}")] = value

    for index, source in enumerate(batch.sources):
        row = DATA_START_ROW + index
        values = {
            "A": source.source_id,
            "B": source.league,
            "C": source.season,
            "D": source.provider,
            "E": source.url,
            "F": to_serial(source.retrieved_utc),
            "G": source.coverage,
            "H": source.status,
            "I": int(source.rows),
            "J": int(source.missing_corners),
            "K": source.attempts,
            "L": source.sha256,
        }
        for letter, value in values.items():
            overrides[(SOURCE_SHEET, f"{letter}{row}")] = value

    for index, record in enumerate(batch.oos_records):
        row = DATA_START_ROW + index
        for offset, value in enumerate(record[:OOS_COLUMNS]):
            overrides[(OOS_SHEET, f"{get_column_letter(offset + 1)}{row}")] = value

    return overrides
