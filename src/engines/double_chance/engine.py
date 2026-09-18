"""API public al motorului Șansă Dublă V2."""

from __future__ import annotations

from dataclasses import dataclass, field
from decimal import ROUND_HALF_UP, Decimal
from functools import lru_cache
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook
from openpyxl.workbook.workbook import Workbook

from config.settings import TEMPLATES_DIR
from src.engines.over05.evaluator import XLException, WorkbookEngine

TEMPLATE_NAME = "11_model_analiza_pariu_sansadubla_optimizat_V2.xlsx"
SELECTIONS: tuple[tuple[str, int], ...] = (("1X", 4), ("X2", 5), ("12", 6))
OFFICIAL_COLUMNS: tuple[str, ...] = (
    "B",
    "D",
    "E",
    "F",
    "H",
    "I",
    "K",
    "L",
    "M",
    "N",
    "O",
    "P",
    "Q",
    "R",
    "S",
    "T",
    "U",
    "V",
    "X",
    "Y",
    "AA",
    "AB",
    "AC",
    "AD",
    "AE",
    "AF",
)
OFFICIAL_CELLS: tuple[str, ...] = tuple(
    f"Double_Chance!{col}{row}" for _, row in SELECTIONS for col in OFFICIAL_COLUMNS
)
# Celule de diagnostic: descompun haircut-ul, draw gate-ul și coeficienții de risc.
# Sunt citite din workbook, nu rescrise, ca auditul să nu dubleze constantele.
DIAGNOSTIC_CELLS: tuple[str, ...] = (
    "Uncertainty_v2!B8",
    "Uncertainty_v2!B9",
    "Uncertainty_v2!B10",
    "Model_1X2!B10",
    "Model_1X2!C10",
    "Model_1X2!D10",
    "P0_Gates_v2!I6",
    "P0_Gates_v2!J6",
    "Parametri!B35",
    "Parametri!B36",
    "Parametri!B37",
    "Parametri!B38",
    "Parametri!B43",
    "Parametri!B44",
    "Parametri!B45",
    "Parametri!B46",
    # Override Protected Strength: input manual (DA/NU) pentru 1X, respectiv X2.
    "Input_Meci!B37",
    "Input_Meci!B38",
)
# Celula de override corespunzătoare fiecărui rând de selecție; 12 nu are override.
OVERRIDE_CELL_BY_ROW: dict[int, str] = {4: "Input_Meci!B37", 5: "Input_Meci!B38"}
# Celule libere pe Input_Meci, după zona de note; nu sunt formule.
PYTHON_RESULT_CELLS: dict[str, str] = {
    "A60": "Selecție",
    "B60": "P_adj (F)",
    "C60": "Pfail_DC (Y)",
    "D60": "P0 final (U)",
    "E60": "Scor (L)",
    "F60": "Nivel (M)",
    "G60": "Verdict (N)",
    "H60": "Eligibil clasament (AA)",
    "I60": "Eligibil DEFENSIV (AE)",
    "J60": "Release (AD)",
}


@dataclass
class RiskBreakdown:
    """Componentele scorului de risc, citite din celulele oficiale.

    Formula V2 rămâne în Excel; aici doar se descompune rezultatul ei, ca un
    FAIL marginal să poată fi distins de unul extrem după plafonarea la 100.
    """

    failure_weight: float | None
    failure_component: float | None
    haircut_weight: float | None
    haircut_component: float | None
    market_direction_surcharge: float
    protected_strength_surcharge: float
    draw_gate_surcharge: float
    away_favorite_surcharge: float
    gate_fail_surcharge: float
    score_before_cap: float | None
    score_final: int | float | None

    def as_dict(self) -> dict[str, Any]:
        return {
            "risk_failure_component": self.failure_component,
            "risk_haircut_component": self.haircut_component,
            "prudent_gate_surcharge": self.market_direction_surcharge
            + self.protected_strength_surcharge,
            "draw_gate_surcharge": self.draw_gate_surcharge,
            "away_favorite_surcharge": self.away_favorite_surcharge,
            "gate_fail_surcharge": self.gate_fail_surcharge,
            "risk_score_before_cap": self.score_before_cap,
            "risk_score_final": self.score_final,
        }


@dataclass
class DrawDiagnostics:
    """Descompunerea Draw Gate-ului pentru 12 (P0_Gates_v2 I6/J6)."""

    model_raw: float | None
    model_plus_haircut: float | None
    market: float | None
    adjusted: float | None


@dataclass
class DoubleChanceSelection:
    """Rezultat oficial pentru o selecție 1X / X2 / 12."""

    code: str
    row: int
    p_model: float | None
    pfail_model: float | None
    pfail_market: float | None
    pfail_dc: float | None
    haircut_base: float | None
    haircut_early: float | None
    haircut_total: float | None
    p_adj: float | None
    p_market: float | None
    edge: float | None
    control: str | None
    market_direction: str | None
    protected_strength: str | None
    protected_strength_override: str | None
    draw_gate: str | None
    p0_final: str | None
    early_season: str | None
    confidence: str | None
    motivation: str | None
    away_favorite_surcharge: float | None
    score: int | float | None
    level: int | float | None
    verdict: str | None
    fragility_gate: str | None
    ranking_eligible: str | None
    defensive_eligible: str | None
    release: str | None
    risk: RiskBreakdown
    draw: DrawDiagnostics | None


@dataclass
class DoubleChanceOfficialResult:
    """Rezultatele oficiale ale celor trei selecții plus celulele evaluate."""

    selections: list[DoubleChanceSelection]
    cells: dict[str, Any] = field(default_factory=dict)

    def by_code(self, code: str) -> DoubleChanceSelection:
        for item in self.selections:
            if item.code == code:
                return item
        raise KeyError(code)

    def summary_recommendation(self) -> str:
        return " | ".join(f"{item.code}: {item.verdict or '—'}" for item in self.selections)

    def payload(self) -> list[dict[str, Any]]:
        rows: list[dict[str, Any]] = []
        for item in self.selections:
            row: dict[str, Any] = {
                "selection": item.code,
                "p_model": item.p_model,
                "p_adj": item.p_adj,
                "p_market": item.p_market,
                "edge": item.edge,
                "pfail_model": item.pfail_model,
                "pfail_market": item.pfail_market,
                "pfail_dc": item.pfail_dc,
                "haircut_base": item.haircut_base,
                "haircut_early": item.haircut_early,
                "haircut_total": item.haircut_total,
                "control": item.control,
                "market_direction": item.market_direction,
                "protected_strength": item.protected_strength,
                "protected_strength_override": item.protected_strength_override,
                "draw_gate": item.draw_gate,
                "p0_final": item.p0_final,
                "early_season": item.early_season,
                "confidence": item.confidence,
                "motivation": item.motivation,
                "score": item.score,
                "level": item.level,
                "verdict": item.verdict,
                "fragility": item.fragility_gate,
                "ranking_eligible": item.ranking_eligible,
                "defensive_eligible": item.defensive_eligible,
                "release": item.release,
                **item.risk.as_dict(),
            }
            if item.draw is not None:
                row.update(
                    {
                        "pdraw_model_raw": item.draw.model_raw,
                        "pdraw_model_plus_haircut": item.draw.model_plus_haircut,
                        "pdraw_market": item.draw.market,
                        "pdraw_adj": item.draw.adjusted,
                    }
                )
            rows.append(row)
        return rows


@lru_cache(maxsize=1)
def load_double_chance_workbook(path: str | None = None) -> Workbook:
    workbook_path = Path(path) if path else TEMPLATES_DIR / TEMPLATE_NAME
    return load_workbook(workbook_path, data_only=False)


def _as_number_or_none(value: Any) -> float | None:
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    return None


def _as_int_if_whole(value: Any) -> int | float | None:
    number = _as_number_or_none(value)
    if number is None:
        return None
    if number == int(number):
        return int(number)
    return number


def _as_text(value: Any) -> str | None:
    if value is None or value == "":
        return None
    if isinstance(value, str):
        return value
    return None


def _parse_addr(addr: str) -> tuple[str, str]:
    sheet, coord = addr.split("!", 1)
    return sheet.strip("'").replace("''", "'"), coord.replace("$", "")


def compute_double_chance(
    inputs: dict[str, Any],
    *,
    evaluate: Iterable[str] | None = None,
    workbook_path: str | Path | None = None,
) -> DoubleChanceOfficialResult:
    """Calculează Șansă Dublă pe șablonul V2, cu override-uri de input.

    Celulele nespecificate rămân cele din șablon, ca în `verifica_excel.ps1`.
    Formulele, ponderile și pragurile nu sunt rescrise.
    """
    wb = load_double_chance_workbook(str(workbook_path) if workbook_path else None)
    overrides: dict[tuple[str, str], Any] = {}
    for addr, value in inputs.items():
        overrides[_parse_addr(addr)] = value

    engine = WorkbookEngine(wb, overrides)
    cells: dict[str, Any] = {}
    wanted = list(dict.fromkeys([*(evaluate or []), *OFFICIAL_CELLS, *DIAGNOSTIC_CELLS]))
    for addr in wanted:
        sheet, coord = _parse_addr(addr)
        try:
            cells[addr] = engine.cell(sheet, coord)
        except (XLException, ZeroDivisionError, ValueError, OverflowError, NotImplementedError) as exc:
            cells[addr] = f"ERROR:{type(exc).__name__}"

    selections: list[DoubleChanceSelection] = []
    for code, row in SELECTIONS:
        def col(name: str, current_row: int = row) -> Any:
            return cells.get(f"Double_Chance!{name}{current_row}")

        score = _as_int_if_whole(col("L"))
        selections.append(
            DoubleChanceSelection(
                code=code,
                row=row,
                p_model=_as_number_or_none(col("B")),
                pfail_model=_as_number_or_none(col("D")),
                pfail_market=_as_number_or_none(col("X")),
                pfail_dc=_as_number_or_none(col("Y")),
                haircut_base=_as_number_or_none(cells.get("Uncertainty_v2!B8")),
                haircut_early=_as_number_or_none(cells.get("Uncertainty_v2!B9")),
                haircut_total=_as_number_or_none(col("E")),
                p_adj=_as_number_or_none(col("F")),
                p_market=_as_number_or_none(col("H")),
                edge=_as_number_or_none(col("I")),
                control=_as_text(col("K")),
                market_direction=_as_text(col("O")),
                protected_strength=_as_text(col("P")),
                protected_strength_override=_as_text(
                    cells.get(OVERRIDE_CELL_BY_ROW.get(row, ""))
                ),
                draw_gate=_as_text(col("Q")),
                p0_final=_as_text(col("U")),
                early_season=_as_text(col("R")),
                confidence=_as_text(col("S")),
                motivation=_as_text(col("V")),
                away_favorite_surcharge=_as_number_or_none(col("T")),
                score=score,
                level=_as_int_if_whole(col("M")),
                verdict=_as_text(col("N")),
                fragility_gate=_as_text(col("AB")),
                ranking_eligible=_as_text(col("AA")),
                defensive_eligible=_as_text(col("AE")),
                release=_as_text(col("AD")),
                risk=_risk_breakdown(cells, row, score),
                draw=_draw_diagnostics(cells) if code == "12" else None,
            )
        )
    return DoubleChanceOfficialResult(selections=selections, cells=cells)


def _risk_breakdown(cells: dict[str, Any], row: int, score: int | float | None) -> RiskBreakdown:
    """Reconstituie componentele lui Double_Chance!L, fără a recalcula verdictul.

    Coeficienții vin din Parametri (B35:B38, B43:B46), nu din literali Python,
    ca decompoziția să nu poată devia de la workbook.
    """
    param = {name: _as_number_or_none(cells.get(f"Parametri!{name}")) for name in
             ("B35", "B36", "B37", "B38", "B43", "B44", "B45", "B46")}
    is_draw_row = row == 6
    failure_weight = param["B37"] if is_draw_row else param["B35"]
    haircut_weight = param["B38"] if is_draw_row else param["B36"]
    pfail_dc = _as_number_or_none(cells.get(f"Double_Chance!Y{row}"))
    haircut = _as_number_or_none(cells.get(f"Double_Chance!E{row}"))
    prudent = param["B44"] or 0.0
    watch = param["B45"] or 0.0
    fail = param["B46"] or 0.0

    failure_component = (
        pfail_dc * failure_weight if pfail_dc is not None and failure_weight is not None else None
    )
    haircut_component = (
        haircut * haircut_weight if haircut is not None and haircut_weight is not None else None
    )
    market_direction = _as_text(cells.get(f"Double_Chance!O{row}"))
    protected = _as_text(cells.get(f"Double_Chance!P{row}"))
    draw_gate = _as_text(cells.get(f"Double_Chance!Q{row}"))
    p0_final = _as_text(cells.get(f"Double_Chance!U{row}"))
    away_surcharge = _as_number_or_none(cells.get(f"Double_Chance!T{row}")) or 0.0

    md_surcharge = 0.0
    ps_surcharge = 0.0
    dg_surcharge = 0.0
    fail_surcharge = 0.0
    if is_draw_row:
        if draw_gate == "PRUDENT":
            dg_surcharge = prudent
        elif draw_gate == "WATCH":
            dg_surcharge = watch
        elif draw_gate == "FAIL":
            dg_surcharge = fail
        away_surcharge = 0.0
    else:
        md_surcharge = prudent if market_direction == "PRUDENT" else 0.0
        ps_surcharge = prudent if protected == "PRUDENT" else 0.0
        fail_surcharge = fail if p0_final == "FAIL" else 0.0
        if row != 5:
            away_surcharge = 0.0

    before_cap: float | None = None
    if failure_component is not None and haircut_component is not None:
        raw = (
            failure_component
            + haircut_component
            + md_surcharge
            + ps_surcharge
            + dg_surcharge
            + away_surcharge
            + fail_surcharge
        )
        before_cap = float(
            Decimal(str(raw)).quantize(Decimal(1), rounding=ROUND_HALF_UP)
        )

    return RiskBreakdown(
        failure_weight=failure_weight,
        failure_component=failure_component,
        haircut_weight=haircut_weight,
        haircut_component=haircut_component,
        market_direction_surcharge=md_surcharge,
        protected_strength_surcharge=ps_surcharge,
        draw_gate_surcharge=dg_surcharge,
        away_favorite_surcharge=away_surcharge,
        gate_fail_surcharge=fail_surcharge,
        score_before_cap=before_cap,
        score_final=score,
    )


def _draw_diagnostics(cells: dict[str, Any]) -> DrawDiagnostics:
    """Separă P(X) model, P(X) model+haircut și P(X) piață pentru selecția 12.

    Draw Gate compară max(model+haircut, piață), în timp ce riscul folosește
    Y=max(model, piață). Diferența este vizibilă abia dacă ambele sunt expuse.
    """
    model_plus_haircut = _as_number_or_none(cells.get("P0_Gates_v2!I6"))
    market = _as_number_or_none(cells.get("P0_Gates_v2!J6"))
    adjusted = None
    if model_plus_haircut is not None and market is not None:
        adjusted = max(model_plus_haircut, market)
    return DrawDiagnostics(
        model_raw=_as_number_or_none(cells.get("Model_1X2!C10")),
        model_plus_haircut=model_plus_haircut,
        market=market,
        adjusted=adjusted,
    )


def python_result_writes(result: DoubleChanceOfficialResult) -> dict[str, object]:
    """Valorile Python de scris pe Input_Meci A60:J63, fără a atinge formule."""
    writes: dict[str, object] = dict(PYTHON_RESULT_CELLS)
    for index, selection in enumerate(result.selections, start=61):
        writes[f"A{index}"] = selection.code
        writes[f"B{index}"] = selection.p_adj
        writes[f"C{index}"] = selection.pfail_dc
        writes[f"D{index}"] = selection.p0_final
        writes[f"E{index}"] = selection.score
        writes[f"F{index}"] = selection.level
        writes[f"G{index}"] = selection.verdict
        writes[f"H{index}"] = selection.ranking_eligible
        writes[f"I{index}"] = selection.defensive_eligible
        writes[f"J{index}"] = selection.release
    return writes
