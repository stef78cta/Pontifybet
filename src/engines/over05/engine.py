"""API public al motorului Over 0.5 V4."""

from __future__ import annotations

from dataclasses import dataclass, field
from functools import lru_cache
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.workbook.workbook import Workbook

from config.settings import TEMPLATES_DIR
from src.engines.over05.evaluator import XLException, WorkbookEngine
from src.engines.over05.inputs import ANALYSIS_SHEET, INPUT_COLUMNS

OFFICIAL_COLUMNS: tuple[str, ...] = ("GL", "GP", "GT", "HE", "HG", "HH", "HK")
TRACE_COLUMNS: tuple[str, ...] = ("GJ", "GM", "GO")
TEMPLATE_NAME = "1_model_analiza_over0.5_optimizat_V4.xlsx"

# Coloane libere, imediat după HK; valorile Python, nu formule Excel.
PYTHON_RESULT_COLUMNS: dict[str, str] = {
    "HL": "P0 recalibrat (GL)",
    "HM": "Over 0.5 raportat (GP)",
    "HN": "Confidence (GT)",
    "HO": "Risk Score (HE)",
    "HP": "G0 model (HG)",
    "HQ": "Nivel risc (HH)",
    "HR": "Recomandare finala (HK)",
}


@dataclass
class Over05OfficialResult:
    """Rezultatele oficiale ale modelului V4 / selectorului v9."""

    p0_recalibrated: float | None
    p_over_reported: float | None
    confidence: float | None
    risk_score: float | None
    g0: str | None
    risk_level: int | float | None
    recommendation: str | None
    cells: dict[str, Any] = field(default_factory=dict)


@lru_cache(maxsize=1)
def load_over05_workbook(path: str | None = None) -> Workbook:
    """Încarcă șablonul o dată; evaluatorul nu mută celulele, doar override-uri."""
    workbook_path = Path(path) if path else TEMPLATES_DIR / TEMPLATE_NAME
    return load_workbook(workbook_path, data_only=False)


def _as_number_or_none(value: Any) -> float | None:
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    return None


def _as_int_if_whole(value: Any) -> int | float | None:
    number = _as_number_or_none(value)
    if number is None:
        return None
    if number == int(number):
        return int(number)
    return number


def compute_over05(
    inputs: dict[str, Any],
    *,
    excel_row: int = 4,
    formula_overrides: dict[str, Any] | None = None,
    workbook_path: str | Path | None = None,
) -> Over05OfficialResult:
    """Calculează rezultatele oficiale pentru un set de inputuri pe un rând.

    Celulele de input nespecificate sunt golite, ca valorile DEMO din șablon
    să nu se scurgă în calcul. Formulele, ponderile și pragurile rămân în workbook.

    @param inputs - mapare coloană Excel → valoare (A, I, EV, …).
    @param excel_row - rândul de evaluat (4–103 în șablon).
    @param formula_overrides - înlocuiri de celule calculate, doar pentru teste
        (ex. injecția HE la pragurile selectorului).
    @param workbook_path - șablon alternativ; implicit templates/.
    @returns {Over05OfficialResult} - GL, GP, GT, HE, HG, HH, HK plus celule de urmă.
    """
    wb = load_over05_workbook(str(workbook_path) if workbook_path else None)
    overrides: dict[tuple[str, str], Any] = {}
    for col in INPUT_COLUMNS:
        overrides[(ANALYSIS_SHEET, f"{col}{excel_row}")] = None
    for col, value in inputs.items():
        overrides[(ANALYSIS_SHEET, f"{col}{excel_row}")] = value
    if formula_overrides:
        for col, value in formula_overrides.items():
            overrides[(ANALYSIS_SHEET, f"{col}{excel_row}")] = value

    engine = WorkbookEngine(wb, overrides)
    cells: dict[str, Any] = {}
    for col in (*OFFICIAL_COLUMNS, *TRACE_COLUMNS):
        try:
            cells[col] = engine.cell(ANALYSIS_SHEET, f"{col}{excel_row}")
        except (XLException, ZeroDivisionError, ValueError, OverflowError, NotImplementedError) as exc:
            cells[col] = f"ERROR:{type(exc).__name__}"

    recommendation = cells.get("HK") if isinstance(cells.get("HK"), str) else None
    if not recommendation:
        recommendation = "WATCH / NO BET"
    return Over05OfficialResult(
        p0_recalibrated=_as_number_or_none(cells.get("GL")),
        p_over_reported=_as_number_or_none(cells.get("GP")),
        confidence=_as_number_or_none(cells.get("GT")),
        risk_score=_as_number_or_none(cells.get("HE")),
        g0=cells.get("HG") if isinstance(cells.get("HG"), str) else None,
        risk_level=_as_int_if_whole(cells.get("HH")),
        recommendation=recommendation,
        cells=cells,
    )


def python_result_values(result: Over05OfficialResult) -> dict[str, object]:
    """Valorile oficiale de scris în HL:HR, vizibile fără recalcul Excel."""
    return {
        "HL": result.p0_recalibrated,
        "HM": result.p_over_reported,
        "HN": result.confidence,
        "HO": result.risk_score,
        "HP": result.g0,
        "HQ": result.risk_level,
        "HR": result.recommendation or "WATCH / NO BET",
    }
