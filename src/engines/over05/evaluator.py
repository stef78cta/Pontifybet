"""Evaluator Excel-compatibil, comun modelelor Over 0.5 V4, Șansă Dublă V2 și Cornere V14.

Reproduce doar funcțiile și convențiile necesare formulelor din workbook-uri.
Nu este un motor Excel general: fiecare extindere este adăugată pentru o
construcție prezentă efectiv în inventarul de formule al unui model.
"""

from __future__ import annotations

import math
import re
from datetime import date, datetime, timedelta
from decimal import Decimal, ROUND_HALF_UP
from functools import lru_cache
from typing import Any

from openpyxl.formula.tokenizer import Tokenizer
from openpyxl.utils.cell import get_column_letter, range_boundaries
from openpyxl.workbook.workbook import Workbook


class XLException(Exception):
    """Eroare de evaluare echivalentă cu o eroare Excel (#VALUE!, #N/A, ciclu)."""


class FormulaSupportError(Exception):
    """Construcție Excel neacoperită de evaluator.

    Separată de `XLException` intenționat: o funcție nesuportată este un defect
    al evaluatorului, nu o eroare de model. `IFERROR` nu o captează, iar motoarele
    o propagă ca eroare tehnică, nu ca verdict sportiv.
    """


# Excel numără zilele de la 1899-12-30 (compatibil cu bug-ul anului 1900 pentru
# orice dată ulterioară lui 1900-03-01, singurul interval folosit de modele).
EXCEL_EPOCH = datetime(1899, 12, 30)


def to_serial(value: datetime | date) -> float:
    """Convertește o dată/dată-oră în serialul zecimal folosit de Excel."""
    moment = value if isinstance(value, datetime) else datetime(value.year, value.month, value.day)
    delta = moment - EXCEL_EPOCH
    return delta.days + delta.seconds / 86400.0 + delta.microseconds / 86_400_000_000.0


def from_serial(serial: float) -> datetime:
    """Inversul lui `to_serial`; folosit de `TEXT` cu formate de dată."""
    return EXCEL_EPOCH + timedelta(days=float(serial))


@lru_cache(maxsize=65536)
def parse_formula(formula: str) -> tuple:
    """Parsează o formulă Excel într-un AST prefixat."""
    tokens = [t for t in Tokenizer(formula).items if t.type != "WHITE-SPACE"]
    index = 0
    prec = {
        "=": 1,
        "<>": 1,
        "<": 1,
        ">": 1,
        "<=": 1,
        ">=": 1,
        "&": 2,
        "+": 3,
        "-": 3,
        "*": 4,
        "/": 4,
        "^": 5,
    }

    def expr(min_prec: int = 0) -> tuple:
        nonlocal index
        token = tokens[index]
        index += 1
        if token.type == "OPERAND":
            if token.subtype == "RANGE":
                node: tuple = ("ref", token.value)
            elif token.subtype == "TEXT":
                node = ("lit", token.value[1:-1].replace('""', '"'))
            elif token.subtype == "NUMBER":
                node = ("lit", float(token.value))
            else:
                node = ("lit", token.value == "TRUE")
        elif token.type == "OPERATOR-PREFIX":
            node = ("unary", token.value, expr(6))
        elif token.type == "PAREN" and token.subtype == "OPEN":
            node = expr()
            assert tokens[index].subtype == "CLOSE"
            index += 1
        elif token.type == "FUNC" and token.subtype == "OPEN":
            args: list[tuple] = []
            while not (tokens[index].type == "FUNC" and tokens[index].subtype == "CLOSE"):
                args.append(expr())
                if tokens[index].type == "SEP":
                    index += 1
                else:
                    break
            assert tokens[index].subtype == "CLOSE"
            index += 1
            node = ("call", token.value[:-1], tuple(args))
        else:
            raise ValueError((token.value, formula))

        while index < len(tokens):
            token = tokens[index]
            if token.type == "OPERATOR-POSTFIX":
                index += 1
                node = ("op", "/", node, ("lit", 100))
                continue
            if token.type != "OPERATOR-INFIX" or prec[token.value] < min_prec:
                break
            index += 1
            node = ("op", token.value, node, expr(prec[token.value] + 1))
        return node

    ast = expr()
    if index != len(tokens):
        raise ValueError((formula, index, len(tokens)))
    return ast


def flatten(value: Any):
    if isinstance(value, list):
        for item in value:
            # Cazul plat este dominant (un interval de celule), deci evităm
            # recursia pentru fiecare element.
            if isinstance(item, list):
                yield from flatten(item)
            else:
                yield item
    else:
        yield value


def as_flat_list(value: Any) -> list[Any]:
    """Listă plată dintr-un interval, reutilizând obiectul când e deja plat."""
    if isinstance(value, list):
        for item in value:
            if isinstance(item, list):
                return list(flatten(value))
        return value
    return [value]


def is_number(value: Any) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def excel_number(value: Any) -> float:
    """Convertește o valoare la număr după aritmetica Excel (blank → 0, TRUE → 1)."""
    if value is None or value == "":
        return 0.0
    if isinstance(value, bool):
        return 1.0 if value else 0.0
    if isinstance(value, (int, float)):
        return float(value)
    raise XLException("VALUE")


def excel_text(value: Any) -> str:
    if value is None:
        return ""
    if is_number(value):
        return str(int(value)) if value == int(value) else str(value)
    return str(value)


def _has_wildcard(pattern: str) -> bool:
    """Spune dacă un criteriu conține `*`/`?` neescapate (Excel escapează cu `~`)."""
    if "*" not in pattern and "?" not in pattern:
        return False
    if "~" not in pattern:
        return True
    index = 0
    while index < len(pattern):
        char = pattern[index]
        if char == "~":
            index += 2
            continue
        if char in "*?":
            return True
        index += 1
    return False


@lru_cache(maxsize=1024)
def _wildcard_regex(pattern: str) -> re.Pattern[str]:
    """Traduce un criteriu Excel cu wildcard în regex ancorat, case-insensitive."""
    parts: list[str] = []
    index = 0
    while index < len(pattern):
        char = pattern[index]
        if char == "~" and index + 1 < len(pattern) and pattern[index + 1] in "*?~":
            parts.append(re.escape(pattern[index + 1]))
            index += 2
            continue
        if char == "*":
            parts.append(".*")
        elif char == "?":
            parts.append(".")
        else:
            parts.append(re.escape(char))
        index += 1
    return re.compile("".join(parts), re.IGNORECASE | re.DOTALL)


def _wildcard_match(value: Any, pattern: str) -> bool:
    """Wildcard Excel: se aplică numai textului, iar șirul gol nu se potrivește cu `*`."""
    if not isinstance(value, str) or value == "":
        return False
    return bool(_wildcard_regex(pattern).fullmatch(value))


def criteria_match(value: Any, criteria: Any) -> bool:
    """Potrivire COUNTIF/COUNTIFS/SUMIFS, fără sensibilitate la majuscule.

    Wildcard-urile `*`/`?` se aplică exclusiv valorilor text: în Excel,
    `COUNTIFS(interval_numeric;"*")` numără doar celulele care conțin text,
    nu numerele formatate ca dată. V14 folosește acest criteriu pentru a
    detecta rândurile native cu dată non-numerică.
    """
    if isinstance(criteria, str):
        # Un criteriu începe cu operator doar dacă primul caracter îl anunță;
        # verificarea scurtă evită șase `startswith` pe fiecare celulă comparată.
        if criteria[:1] in ("<", ">", "="):
            for op in (">=", "<=", "<>", ">", "<", "="):
                if criteria.startswith(op):
                    rest = criteria[len(op) :]
                    try:
                        right: Any = float(rest)
                    except ValueError:
                        right = rest
                    if is_number(right) and not is_number(value):
                        return False
                    if isinstance(right, str) and _has_wildcard(right):
                        matched = _wildcard_match(value, right)
                        return not matched if op == "<>" else matched
                    return excel_compare(value, right, op)
        if _has_wildcard(criteria):
            return _wildcard_match(value, criteria)
        return excel_compare(value, criteria, "=")
    if is_number(criteria):
        return is_number(value) and float(value) == float(criteria)
    return excel_compare(value, criteria, "=")


def excel_compare(left: Any, right: Any, op: str) -> bool:
    """Comparație Excel: blank se aliniază la tipul celuilalt operand, iar la
    tipuri diferite numerele preced textul (ordinea de sortare Excel)."""
    if left is None:
        left = "" if isinstance(right, str) else 0
    if right is None:
        right = "" if isinstance(left, str) else 0
    left_text = isinstance(left, str)
    right_text = isinstance(right, str)
    if left_text:
        if right_text:
            first: Any = left.casefold()
            second: Any = right.casefold()
        else:
            first, second = 1, 0
    elif right_text:
        first, second = 0, 1
    else:
        first, second = left, right
    if op == "=":
        return first == second
    if op == "<>":
        return first != second
    if op == "<":
        return first < second
    if op == ">":
        return first > second
    if op == "<=":
        return first <= second
    if op == ">=":
        return first >= second
    raise FormulaSupportError(f"operator {op}")


def _sumproduct(arrays: list[Any]) -> float:
    """SUMPRODUCT: sumă a produselor poziționale, valorile non-numerice fiind 0."""
    sized = [item for item in arrays if isinstance(item, list)]
    if not sized:
        product = 1.0
        for item in arrays:
            product *= excel_number(item)
        return product
    length = len(sized[0])
    if any(len(item) != length for item in sized):
        raise XLException("VALUE")
    total = 0.0
    for index in range(length):
        product = 1.0
        for item in arrays:
            value = item[index] if isinstance(item, list) else item
            if isinstance(value, bool):
                product *= 1.0 if value else 0.0
            elif is_number(value):
                product *= float(value)
            else:
                product = 0.0
                break
        total += product
    return total


def _filter(array: Any, include: Any) -> list[Any]:
    """`_xlfn._xlws.FILTER` fără argument `if_empty`: eroare #CALC! pe rezultat vid."""
    items = array if isinstance(array, list) else [array]
    mask = include if isinstance(include, list) else [include] * len(items)
    if len(mask) != len(items):
        raise XLException("VALUE")
    kept = [value for value, flag in zip(items, mask) if excel_number(flag) != 0]
    if not kept:
        raise XLException("CALC")
    return kept


def excel_text_format(value: Any, fmt: Any) -> str:
    """`TEXT` pentru formatele folosite de modele: date și zecimale fixe.

    Formatele de dată sunt recunoscute prin absența cifrelor din cod (`yyyy-mm-dd`,
    `yyyymmdd`), iar valoarea este interpretată ca serial Excel.
    """
    if not isinstance(fmt, str):
        raise XLException("VALUE")
    if any(char in fmt for char in "ymd") and not any(char.isdigit() for char in fmt):
        moment = from_serial(excel_number(value))
        rendered = fmt.replace("yyyy", f"{moment.year:04d}")
        rendered = rendered.replace("mm", f"{moment.month:02d}")
        return rendered.replace("dd", f"{moment.day:02d}")
    digits = len(fmt.split(".")[1].replace("%", "").replace("x", "")) if "." in fmt else 0
    scaled = excel_number(value) * (100 if "%" in fmt else 1)
    quantized = Decimal(str(scaled)).quantize(
        Decimal(1).scaleb(-digits), rounding=ROUND_HALF_UP
    )
    suffix = "%" if "%" in fmt else ("x" if fmt.endswith("x") else "")
    return f"{quantized:.{digits}f}{suffix}"


def _scalar_op(op: str, left: Any, right: Any) -> Any:
    """Aplică un operator Excel pe două scalari."""
    if op in {"=", "<>", "<", ">", "<=", ">="}:
        return excel_compare(left, right, op)
    if op == "&":
        return excel_text(left) + excel_text(right)
    x = excel_number(left)
    y = excel_number(right)
    if op == "+":
        return x + y
    if op == "-":
        return x - y
    if op == "*":
        return x * y
    if op == "/":
        return x / y
    if op == "^":
        return x**y
    raise FormulaSupportError(f"operator {op}")


def apply_op(op: str, left: Any, right: Any) -> Any:
    """Aplică un operator Excel, cu semantică de matrice element cu element.

    Expresiile V14 de tip `(interval=valoare)*(interval<valoare)` produc matrici
    de booleeni pe care Excel le înmulțește pozițional. Scalarii se difuzează
    (broadcast) peste matrice, iar două matrici trebuie să aibă aceeași lungime.
    """
    left_is_array = isinstance(left, list)
    right_is_array = isinstance(right, list)
    if not left_is_array and not right_is_array:
        return _scalar_op(op, left, right)
    if left_is_array and right_is_array:
        if len(left) != len(right):
            raise XLException("VALUE")
        return [_scalar_op(op, a, b) for a, b in zip(left, right)]
    if left_is_array:
        return [_scalar_op(op, a, right) for a in left]
    return [_scalar_op(op, left, b) for b in right]


#: Funcții care ignoră non-numericele și lucrează pe toate argumentele deodată.
_NUMERIC_AGGREGATES = frozenset({"COUNT", "SUM", "AVERAGE", "MIN", "MAX"})


class WorkbookEngine:
    """Evaluează formulele unui workbook, cu override-uri pe celule de input."""

    def __init__(
        self,
        workbook: Workbook,
        overrides: dict[tuple[str, str], Any] | None = None,
        *,
        dates_as_serial: bool = False,
    ) -> None:
        self.wb = workbook
        self.overrides = overrides or {}
        self.memo: dict[tuple[str, str], Any] = {}
        self.visiting: set[tuple[str, str]] = set()
        # Cornere V14 compară și formatează date prin `INT`, `TEXT` și `ISNUMBER`,
        # care în Excel operează pe serialul numeric. openpyxl întoarce `datetime`,
        # deci modelele care depind de aritmetica pe date cer conversia la serial.
        self.dates_as_serial = dates_as_serial
        self._range_memo: dict[tuple[str, str], list[Any]] = {}
        self._tables: dict[str, tuple[str, str]] | None = None

    def cell(self, sheet: str, coord: str) -> Any:
        coord = coord.replace("$", "")
        key = (sheet, coord)
        if key in self.overrides:
            return self.overrides[key]
        if key in self.memo:
            return self.memo[key]
        if key in self.visiting:
            raise XLException("CYCLE")
        raw = self.wb[sheet][coord].value
        if not (isinstance(raw, str) and raw.startswith("=")):
            if self.dates_as_serial and isinstance(raw, (datetime, date)):
                raw = to_serial(raw)
            self.memo[key] = raw
            return raw
        self.visiting.add(key)
        try:
            result = self.eval_ast(parse_formula(raw), sheet)
        finally:
            self.visiting.remove(key)
        self.memo[key] = result
        return result

    def _table_column_ref(self, ref: str) -> tuple[str, str] | None:
        """Rezolvă o referință structurată `Tabel[Coloană]` la `(foaie, interval)`."""
        match = re.fullmatch(r"([A-Za-z_][\w.]*)\[([^\[\]]+)\]", ref.strip())
        if match is None:
            return None
        if self._tables is None:
            registry: dict[str, tuple[str, str]] = {}
            for worksheet in self.wb.worksheets:
                for name in worksheet.tables:
                    registry[name.casefold()] = (worksheet.title, name)
            self._tables = registry
        entry = self._tables.get(match.group(1).casefold())
        if entry is None:
            return None
        sheet_name, table_name = entry
        table = self.wb[sheet_name].tables[table_name]
        min_col, min_row, max_col, max_row = range_boundaries(table.ref)
        columns = [column.name for column in table.tableColumns]
        wanted = match.group(2).strip().casefold()
        try:
            offset = next(i for i, name in enumerate(columns) if name.casefold() == wanted)
        except StopIteration:
            raise FormulaSupportError(f"coloană inexistentă: {ref}") from None
        letter = get_column_letter(min_col + offset)
        # Rândul de header (primul din `ref`) nu face parte din corpul tabelului.
        return sheet_name, f"{letter}{min_row + 1}:{letter}{max_row}"

    def ref(self, ref: str, sheet: str) -> Any:
        structured = self._table_column_ref(ref)
        if structured is not None:
            sheet, ref = structured
        elif "!" in ref:
            sheet, ref = ref.rsplit("!", 1)
            sheet = sheet.strip("'").replace("''", "'")
        ref = ref.replace("$", "")
        if ":" not in ref:
            return self.cell(sheet, ref)
        key = (sheet, ref)
        cached = self._range_memo.get(key)
        if cached is not None:
            return cached
        min_col, min_row, max_col, max_row = range_boundaries(ref)
        values = [
            self.cell(sheet, f"{get_column_letter(col)}{row}")
            for row in range(min_row, max_row + 1)
            for col in range(min_col, max_col + 1)
        ]
        self._range_memo[key] = values
        return values

    def ref_shape(self, ref: str, sheet: str) -> tuple[int, int]:
        """Numărul de rânduri și coloane al unei referințe, pentru ROWS/COLUMNS."""
        structured = self._table_column_ref(ref)
        if structured is not None:
            sheet, ref = structured
        elif "!" in ref:
            sheet, ref = ref.rsplit("!", 1)
        ref = ref.replace("$", "")
        min_col, min_row, max_col, max_row = range_boundaries(ref)
        return max_row - min_row + 1, max_col - min_col + 1

    def eval_ast(self, ast: tuple, sheet: str) -> Any:
        kind = ast[0]
        if kind == "lit":
            return ast[1]
        if kind == "ref":
            return self.ref(ast[1], sheet)
        if kind == "unary":
            operand = self.eval_ast(ast[2], sheet)
            sign = 1 if ast[1] == "+" else -1
            if isinstance(operand, list):
                return [excel_number(item) * sign for item in operand]
            return excel_number(operand) * sign
        if kind == "op":
            return apply_op(
                ast[1],
                self.eval_ast(ast[2], sheet),
                self.eval_ast(ast[3], sheet),
            )

        name, args = ast[1], ast[2]
        if name == "IF":
            return self.eval_ast(args[1] if self.eval_ast(args[0], sheet) else args[2], sheet)
        if name == "IFERROR":
            try:
                return self.eval_ast(args[0], sheet)
            except (XLException, ZeroDivisionError, ValueError, OverflowError):
                return self.eval_ast(args[1], sheet)
        if name == "ROWS":
            if args[0][0] == "ref":
                return self.ref_shape(args[0][1], sheet)[0]
            value = self.eval_ast(args[0], sheet)
            return len(value) if isinstance(value, list) else 1

        values = [self.eval_ast(arg, sheet) for arg in args]
        # Agregarea numerică este cerută de puține funcții, dar ar parcurge
        # intervale mari la fiecare apel, deci se calculează la cerere.
        if name in _NUMERIC_AGGREGATES:
            numbers = [
                item
                for item in flatten(values)
                if isinstance(item, (int, float)) and not isinstance(item, bool)
            ]
            if name == "COUNT":
                return len(numbers)
            if name == "SUM":
                return sum(numbers)
            if name == "AVERAGE":
                return sum(numbers) / len(numbers)
            if name == "MIN":
                return min(numbers) if numbers else 0
            return max(numbers) if numbers else 0
        if name == "COUNTA":
            return sum(1 for item in flatten(values) if item is not None and item != "")
        if name == "COUNTBLANK":
            return sum(1 for item in flatten(values) if item is None or item == "")
        if name == "COUNTIF":
            criteria = values[1]
            return sum(1 for item in as_flat_list(values[0]) if criteria_match(item, criteria))
        if name == "COUNTIFS":
            pairs = [
                (as_flat_list(values[i]), values[i + 1]) for i in range(0, len(values) - 1, 2)
            ]
            total = 0
            for idx in range(len(pairs[0][0])):
                for cells, criteria in pairs:
                    if not criteria_match(cells[idx], criteria):
                        break
                else:
                    total += 1
            return total
        if name == "SUMIFS":
            summed = as_flat_list(values[0])
            pairs = [
                (as_flat_list(values[i]), values[i + 1]) for i in range(1, len(values) - 1, 2)
            ]
            total_sum = 0.0
            for idx, item in enumerate(summed):
                for cells, criteria in pairs:
                    if not criteria_match(cells[idx], criteria):
                        break
                else:
                    if isinstance(item, (int, float)) and not isinstance(item, bool):
                        total_sum += float(item)
            return total_sum
        if name == "ISNUMBER":
            return is_number(values[0])
        if name == "NOT":
            flag = values[0]
            if isinstance(flag, bool):
                return not flag
            if is_number(flag):
                return float(flag) == 0.0
            raise XLException("VALUE")
        if name == "MOD":
            numerator = excel_number(values[0])
            denominator = excel_number(values[1])
            return numerator - denominator * math.floor(numerator / denominator)
        if name == "AND":
            return all(values)
        if name == "OR":
            return any(values)
        if name == "ABS":
            return abs(excel_number(values[0]))
        if name == "SQRT":
            return math.sqrt(excel_number(values[0]))
        if name == "EXP":
            return math.exp(excel_number(values[0]))
        if name == "LN":
            return math.log(excel_number(values[0]))
        if name == "POWER":
            return excel_number(values[0]) ** excel_number(values[1])
        if name == "ROUND":
            return float(
                Decimal(str(values[0])).quantize(
                    Decimal(1).scaleb(-int(values[1])),
                    rounding=ROUND_HALF_UP,
                )
            )
        if name == "MATCH":
            if len(values) > 2 and excel_number(values[2]) != 0:
                raise FormulaSupportError("MATCH cu potrivire aproximativă")
            for i, item in enumerate(flatten(values[1]), 1):
                if excel_compare(values[0], item, "="):
                    return i
            raise XLException("NA")
        if name == "INDEX":
            items = values[0] if isinstance(values[0], list) else [values[0]]
            position = int(excel_number(values[1]))
            if position < 1 or position > len(items):
                raise XLException("REF")
            return items[position - 1]
        if name == "INT":
            return float(math.floor(excel_number(values[0])))
        if name == "LEFT":
            count = int(excel_number(values[1])) if len(values) > 1 else 1
            if count < 0:
                raise XLException("VALUE")
            return excel_text(values[0])[:count]
        if name == "SUMPRODUCT":
            return _sumproduct(values)
        if name == "_xlfn._xlws.FILTER":
            return _filter(values[0], values[1])
        if name == "LARGE":
            # Doar primul argument formează mulțimea; `k` nu intră în ea.
            pool = sorted(item for item in flatten(values[0]) if is_number(item))
            pool.reverse()
            position = int(excel_number(values[1]))
            if position < 1 or position > len(pool):
                raise XLException("NUM")
            return pool[position - 1]
        if name == "TEXT":
            return excel_text_format(values[0], values[1])
        raise FormulaSupportError(name)
