"""Generator Excel: copiere șablon + scriere controlată + ZIP."""

from __future__ import annotations

import json
import shutil
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from config.settings import OUTPUTS_DIR, TEMPLATES_DIR, ensure_runtime_dirs
from src.excel.integrity import (
    IntegrityError,
    WorkbookFingerprint,
    assert_whitelist_write,
    compare_fingerprints,
    fingerprint_workbook,
    get_template_baseline,
)


def copy_template(template_name: str, dest: Path) -> Path:
    """Copiază șablonul read-only într-un fișier nou."""
    src = TEMPLATES_DIR / template_name
    if not src.exists():
        raise FileNotFoundError(f"Șablon lipsă: {src}")
    dest.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(src, dest)
    return dest


def unix_to_excel_serial(ts: int | float) -> float:
    """Unix timestamp → serial Excel (zile de la 1899-12-30)."""
    epoch = datetime(1970, 1, 1, tzinfo=timezone.utc)
    dt = datetime.fromtimestamp(float(ts), tz=timezone.utc)
    delta = dt - epoch
    # Excel epoch 1899-12-30; Unix epoch = 25569 zile
    return 25569 + delta.total_seconds() / 86400.0


class SafeWorkbookWriter:
    """Scrie doar pe whitelist; verifică că nu atinge formule."""

    def __init__(
        self,
        path: Path,
        whitelist: set[str],
        *,
        template_baseline: WorkbookFingerprint | None = None,
    ) -> None:
        self.path = path
        self.whitelist = whitelist
        self.before = template_baseline or fingerprint_workbook(path)
        self.wb = load_workbook(path)
        self._writes: list[str] = []

    def write(self, sheet: str, coordinate: str, value: Any) -> None:
        assert_whitelist_write(sheet, coordinate, self.whitelist)
        if sheet not in self.wb.sheetnames:
            raise IntegrityError(f"EXPORT BLOCAT. Motiv: foaia {sheet} nu există.")
        ws = self.wb[sheet]
        cell = ws[coordinate]
        if isinstance(cell.value, str) and cell.value.startswith("="):
            raise IntegrityError(
                f"EXPORT BLOCAT. Motiv: formula {sheet}!{coordinate} a fost modificată."
            )
        cell.value = value
        self._writes.append(f"{sheet}!{coordinate}")

    def save(self) -> Path:
        self.wb.save(self.path)
        self.wb.close()
        after = fingerprint_workbook(self.path)
        compare_fingerprints(self.before, after)
        return self.path


def make_run_dir(run_id: str | None = None) -> Path:
    ensure_runtime_dirs()
    rid = run_id or datetime.now().strftime("%Y%m%d_%H%M%S")
    path = OUTPUTS_DIR / rid
    path.mkdir(parents=True, exist_ok=True)
    return path


def write_validation_report(run_dir: Path, report: dict[str, Any]) -> Path:
    path = run_dir / "validation_report.json"
    path.write_text(json.dumps(report, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    # CSV scurt
    csv_path = run_dir / "validation_report.csv"
    lines = ["match_id,model_id,indicator,g_level,status,blocking,reason"]
    for issue in report.get("issues", []):
        reason = str(issue.get("reason", "")).replace('"', "'")
        lines.append(
            f"{issue.get('match_id')},{issue.get('model_id')},{issue.get('indicator')},"
            f"{issue.get('g_level')},{issue.get('status')},{issue.get('blocking')},\"{reason}\""
        )
    csv_path.write_text("\n".join(lines), encoding="utf-8")
    return path


def write_over05_results_csv(run_dir: Path, rows: list[dict[str, Any]]) -> Path | None:
    """Scrie verdicturile Over 0.5 calculate în Python, vizibile în ZIP.

    Coloanele `ora` și `liga` stau înaintea `echipe`, ca în tabelele din UI.
    """
    official = [r for r in rows if r.get("model_id") == "over05"]
    if not official:
        return None
    path = run_dir / "over05_rezultate.csv"
    lines = [
        "ora,liga,match_id,echipe,recomandare_HK,nivel_HH,g0_HG,p_over_GP,p0_GL,confidence_GT,risk_score_HE"
    ]
    for row in official:
        rec = str(row.get("recommendation") or "").replace('"', "'")
        liga = str(row.get("liga") or "").replace('"', "'")
        lines.append(
            f"\"{row.get('ora_bucuresti') or ''}\",\"{liga}\","
            f"{row.get('match_id')},\"{row.get('echipe','')}\",\"{rec}\","
            f"{row.get('risk_level') or ''},{row.get('model_g0') or ''},"
            f"{row.get('p_over') or ''},{row.get('p0_recalibrated') or ''},"
            f"{row.get('confidence') or ''},{row.get('risk_score') or ''}"
        )
    path.write_text("\n".join(lines), encoding="utf-8")
    return path


def write_double_chance_results_csv(run_dir: Path, rows: list[dict[str, Any]]) -> Path | None:
    """Scrie verdicturile Șansă Dublă calculate în Python, câte un rând pe selecție."""
    official = [r for r in rows if r.get("model_id") == "double_chance"]
    if not official:
        return None
    path = run_dir / "double_chance_rezultate.csv"
    numeric_fields = (
        "level",
        "score",
        "risk_score_before_cap",
        "p_model",
        "p_adj",
        "p_market",
        "pfail_model",
        "pfail_market",
        "pfail_dc",
        "haircut_base",
        "haircut_early",
        "haircut_total",
        "p_scoreline",
        "p_strength",
        "p_market_component",
        "contribution_scoreline",
        "contribution_strength",
        "contribution_market",
        "risk_failure_component",
        "risk_haircut_component",
        "prudent_gate_surcharge",
        "draw_gate_surcharge",
        "away_favorite_surcharge",
        "gate_fail_surcharge",
        "fragility_level_surcharge",
        "pdraw_model_raw",
        "pdraw_model_plus_haircut",
        "pdraw_market",
        "pdraw_adj",
    )
    # Câmpuri de context, identice pentru cele trei selecții ale unui meci.
    context_fields = (
        "dc_sample_status",
        "dc_prior_status",
        "dc_prior_source",
        "dc_prior_method",
        "dc_shrinkage_applied",
        "dc_sample_n_home",
        "dc_sample_n_away",
        "dc_prior_n_home",
        "dc_prior_n_away",
        "dc_weight_current",
        "dc_weight_prior",
        "dc_rate_before_shrinkage",
        "dc_rate_after_shrinkage",
        "dc_lambda_home_before",
        "dc_lambda_home_after",
        "dc_lambda_away_before",
        "dc_lambda_away_after",
    )
    text_fields = (
        "market_direction",
        "protected_strength",
        "protected_strength_override",
        "draw_gate",
        "p0_final",
        "confidence",
        "release",
        "ranking_eligible",
        "defensive_eligible",
    )
    header = [
        "ora",
        "liga",
        "match_id",
        "echipe",
        "selectie",
        "verdict_N",
        "nivel_M",
        "scor_L",
        "scor_brut",
        "p_model_B",
        "p_adj_F",
        "p_market_H",
        "pfail_model_D",
        "pfail_market_X",
        "pfail_dc_Y",
        "haircut_baza",
        "haircut_early",
        "haircut_total_E",
        "p_scoreline",
        "p_strength",
        "p_market_devig",
        "contrib_scoreline",
        "contrib_strength",
        "contrib_market",
        "risk_failure",
        "risk_haircut",
        "surcharge_prudent",
        "surcharge_draw",
        "surcharge_away_fav",
        "surcharge_fail",
        "surcharge_fragility_nivel",
        "pdraw_model_raw",
        "pdraw_model_dupa_haircut",
        "pdraw_market",
        "pdraw_adjusted",
        "market_direction_O",
        "protected_strength_P",
        "ps_override_B37_B38",
        "draw_gate_Q",
        "p0_final_U",
        "confidence_S",
        "release_AD",
        "eligibil_AA",
        "defensiv_AE",
        "esantion_C9",
        "prior_C10",
        "prior_sursa_D10",
        "prior_metoda_I10",
        "shrinkage_aplicat",
        "n_current_home",
        "n_current_away",
        "n_prior_home",
        "n_prior_away",
        "weight_current",
        "weight_prior",
        "rata_inainte_shrinkage",
        "rata_dupa_shrinkage",
        "lambda_home_inainte",
        "lambda_home_dupa",
        "lambda_away_inainte",
        "lambda_away_dupa",
    ]
    lines = [",".join(header)]

    def number(value: object) -> str:
        return "" if value is None or value == "" else str(value)

    def quoted(value: object) -> str:
        return '"' + str(value or "").replace('"', "'") + '"'

    for row in official:
        prefix = [
            quoted(row.get("ora_bucuresti")),
            quoted(row.get("liga")),
            str(row.get("match_id") or ""),
            quoted(row.get("echipe")),
        ]
        context = [
            quoted(row.get(name))
            if isinstance(row.get(name), str) or row.get(name) is None
            else number(row.get(name))
            for name in context_fields
        ]
        selections = row.get("dc_selections") or []
        if not selections:
            middle = [quoted(""), quoted(row.get("recommendation")), number(row.get("risk_level"))]
            padding = [""] * (len(header) - len(prefix) - len(middle) - len(context))
            lines.append(",".join([*prefix, *middle, *padding, *context]))
            continue
        for sel in selections:
            values = [
                str(sel.get("selection") or ""),
                quoted(sel.get("verdict")),
                *[number(sel.get(name)) for name in numeric_fields],
                *[quoted(sel.get(name)) for name in text_fields],
            ]
            lines.append(",".join([*prefix, *values, *context]))
    path.write_text("\n".join(lines), encoding="utf-8")
    return path


def zip_outputs(run_dir: Path, zip_name: str = "pontifybet_export.zip") -> Path:
    zip_path = run_dir / zip_name
    with zipfile.ZipFile(zip_path, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        for file in run_dir.iterdir():
            if file.is_file() and file.name != zip_name:
                zf.write(file, arcname=file.name)
    return zip_path


def cleanup_run_dir(run_dir: Path) -> None:
    if run_dir.exists() and run_dir.is_dir():
        shutil.rmtree(run_dir, ignore_errors=True)
