"""Ingestia verdictelor Cornere V14 din workbook-ul recalculat de Microsoft Excel.

De ce există acest modul: verdictul Cornere nu este disponibil în Python. Pipeline-ul
scrie doar inputuri în copia șablonului, iar `Analiza_Linii` (Verdict_FINAL,
Risk_Level_FINAL, P_FINAL, …) se calculează abia la deschiderea fișierului în
Microsoft Excel. Alternativa — rescrierea formulelor în Python — ar produce un al
doilea motor de analiză și ar putea divergea silențios de modelul oficial. Așa că
Excel rămâne source of truth, iar aici doar *citim* rezultatul lui.

`openpyxl` nu este folosit ca motor de formule: fișierul este deschis cu
`data_only=True`, adică se citesc exclusiv valorile pe care Excel le-a salvat în
cache. Un workbook necalculat de Excel este respins, nu „reparat”.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from config.settings import TEMPLATES_DIR
from src.adapters.base import load_registry
from src.excel.integrity import (
    IntegrityError,
    compare_fingerprints,
    fingerprint_workbook,
    get_template_baseline,
)
from src.history import repository as repo
from src.history.db import open_db
from src.history.models import Market, PredictionKey, SnapshotStatus
from src.history.snapshots import CORNERS_RESULT_SHEET, model_version

# Whitelist de citire: exclusiv coloanele de rezultat ale modelului, pe foaia de
# rezultate. Nimic din `Config_v6`, `Motor_Meciuri` sau `Distributie` nu este citit
# sau interpretat aici.
READ_COLUMNS: dict[str, str] = {
    "B": "match_id",
    "G": "type",
    "H": "k",
    "J": "hard_gate",
    "S": "failure_mode",
    "T": "risk_score",
    "V": "confidence",
    "Z": "risk_level",
    "AA": "recommendation",
    "AB": "p_adjusted",
    "AH": "motivation",
}
READ_CELL_WHITELIST: frozenset[str] = frozenset(
    f"{CORNERS_RESULT_SHEET}!{col}" for col in READ_COLUMNS
)


def _corners_template_path() -> Path:
    return TEMPLATES_DIR / load_registry()["models"]["corners"]["template"]


def assert_matches_template(path: Path) -> None:
    """Blochează importul dacă workbook-ul nu mai corespunde șablonului.

    Reutilizează mecanismul de integritate existent (foi + formule), ca un fișier
    editat manual să nu poată injecta valori în istoric.
    """
    baseline = get_template_baseline(_corners_template_path())
    try:
        compare_fingerprints(baseline, fingerprint_workbook(path))
    except IntegrityError as exc:
        raise IntegrityError(exc.message.replace("EXPORT BLOCAT", "IMPORT BLOCAT")) from exc


def _as_float(value: Any) -> float | None:
    if isinstance(value, bool) or value is None or value == "":
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _as_text(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def read_recalculated_lines(path: Path) -> list[dict[str, Any]]:
    """Citește liniile Cornere din workbook-ul recalculat.

    @param path - copia exportată, deschisă și salvată de Microsoft Excel.
    @returns - un dict pe linie, cu valorile finale ale modelului.
    @throws {IntegrityError} - workbook necorespunzător șablonului sau nerecalculat.
    """
    if not Path(path).exists():
        raise IntegrityError(f"IMPORT BLOCAT. Motiv: fișierul {path} nu există.")
    assert_matches_template(Path(path))

    wb = load_workbook(path, data_only=True)
    try:
        if CORNERS_RESULT_SHEET not in wb.sheetnames:
            raise IntegrityError(
                f"IMPORT BLOCAT. Motiv: foaia {CORNERS_RESULT_SHEET} lipsește."
            )
        ws = wb[CORNERS_RESULT_SHEET]
        rows: list[dict[str, Any]] = []
        saw_formula_cache = False
        for row in range(6, ws.max_row + 1):
            raw = {name: ws[f"{col}{row}"].value for col, name in READ_COLUMNS.items()}
            match_id = _as_text(raw["match_id"])
            if raw["match_id"] is not None:
                saw_formula_cache = True
            if not match_id:
                continue
            kind = str(raw["type"] or "").strip().upper()
            k = _as_float(raw["k"])
            if kind not in {"OVER", "UNDER"} or k is None:
                continue
            rows.append(
                {
                    "provider_match_id": match_id,
                    "market": (
                        Market.CORNERS_OVER.value
                        if kind == "OVER"
                        else Market.CORNERS_UNDER.value
                    ),
                    "line": k + 0.5,
                    "recommendation": _as_text(raw["recommendation"]),
                    "p_adjusted": _as_float(raw["p_adjusted"]),
                    "failure_mode": _as_float(raw["failure_mode"]),
                    "risk_score": _as_float(raw["risk_score"]),
                    "risk_level": _as_float(raw["risk_level"]),
                    "confidence": _as_float(raw["confidence"]),
                    "hard_gate": _as_text(raw["hard_gate"]),
                    "motivation": _as_text(raw["motivation"]),
                }
            )
        if not saw_formula_cache:
            raise IntegrityError(
                "IMPORT BLOCAT. Motiv: workbook-ul nu a fost recalculat de "
                "Microsoft Excel (celulele de rezultat nu au valori salvate)."
            )
        return rows
    finally:
        wb.close()


def ingest_corners_workbook(path: Path | str, *, db_path: str | None = None) -> dict[str, Any]:
    """Îngheață predicțiile Cornere pe baza workbook-ului recalculat.

    Atinge exclusiv predicțiile aflate în `PENDING_EXCEL_RECALC`. O linie deja
    `FROZEN_PREMATCH` este ignorată — forecastul publicat nu se rescrie.
    """
    lines = read_recalculated_lines(Path(path))
    version = model_version("corners")
    frozen = 0
    skipped_frozen = 0
    unknown = 0
    with open_db(db_path) as conn:
        for line in lines:
            key = PredictionKey(
                provider_match_id=line["provider_match_id"],
                model_key="corners",
                model_version=version,
                market=line["market"],
                line=line["line"],
            )
            existing = repo.find_prediction(conn, key)
            if existing is None:
                unknown += 1
                continue
            if existing["snapshot_status"] == SnapshotStatus.FROZEN_PREMATCH.value:
                skipped_frozen += 1
                continue
            values = {
                name: line[name]
                for name in (
                    "recommendation",
                    "p_adjusted",
                    "failure_mode",
                    "risk_score",
                    "risk_level",
                    "confidence",
                    "hard_gate",
                    "motivation",
                )
            }
            if repo.freeze_prediction(conn, int(existing["id"]), values=values):
                frozen += 1
    return {
        "lines_read": len(lines),
        "frozen": frozen,
        "already_frozen": skipped_frozen,
        "unknown_predictions": unknown,
    }
