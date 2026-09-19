"""Paritate motor Cornere V14 față de cele 42 de selecții din pachetul tehnic.

Fixture-urile sunt cele trei exemple REPLAY livrate cu workbook-ul: inputurile
originale (`fixture_input_original.json`) și rezultatele așteptate pentru toate
cele 40 de coloane ale `LinesTable` (`expected_42_selectii.csv`).
"""

from __future__ import annotations

import csv
import json
import math
from datetime import datetime
from functools import lru_cache
from hashlib import sha256
from pathlib import Path

import pytest

from config.settings import TEMPLATES_DIR
from src.engines.corners import (
    CORNERS_LINES,
    CornersBatchInput,
    CornersMatchInput,
    NativeHistoryRow,
    NativeSource,
    compute_corners_batch,
    rank_top_live,
)
from src.engines.corners.engine import load_corners_workbook
from src.engines.corners.inputs import CONTEXT_FIELDS

ROOT = Path(__file__).resolve().parent.parent
PACK = ROOT / "docs" / "Pachet_tehnic_Cornere_Multiline_V14"
TEMPLATE = TEMPLATES_DIR / "6_model_analiza_cornere_multiline_optimizat_V14.xlsx"

#: Toleranța cerută de pachetul tehnic pentru comparațiile numerice.
ABS_TOL = 1e-10
REL_TOL = 1e-12

#: Coloane textuale în care pachetul a normalizat cratima lungă („–” → „-”).
DASH_NORMALIZED = {"Match"}

#: Coloane a căror valoare este produsă de `Dashboard` pentru clasamentul global.
CROSS_SLOT_COLUMNS = {"Rank_LIVE"}


def _parse_dt(value: str) -> datetime:
    return datetime.fromisoformat(value)


@lru_cache(maxsize=1)
def _cached_workbook():
    """Valorile memorate în șablon, folosite ca referință pentru intermediari."""
    from openpyxl import load_workbook

    return load_workbook(TEMPLATE, data_only=True)


def _serial_to_datetime(serial: float) -> datetime:
    """Convertește un serial Excel în datetime (epoca 1899-12-30, ca openpyxl)."""
    from datetime import timedelta

    return datetime(1899, 12, 30) + timedelta(days=float(serial))


@lru_cache(maxsize=1)
def _fixture() -> dict:
    return json.loads((PACK / "fixture_input_original.json").read_text(encoding="utf-8"))


@lru_cache(maxsize=1)
def _expected() -> list[dict[str, str]]:
    with (PACK / "expected_42_selectii.csv").open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


@lru_cache(maxsize=1)
def _batch() -> CornersBatchInput:
    data = _fixture()
    matches = tuple(
        CornersMatchInput(
            match_id=row["Match_ID"],
            league=row["League"],
            season=row["Season"],
            home=row["Home"],
            away=row["Away"],
            kickoff_utc=_parse_dt(row["Kickoff_UTC"]),
            cutoff_utc=_parse_dt(row["Cutoff_UTC"]),
            mode=row["Mode"],
            sourcing_decision=row.get("Sourcing_Decision") or "",
            evidence_or_attempts=row.get("Evidence_or_Attempts") or "",
            prior_season=row.get("Prior_Season") or "",
            prior_allowed=row.get("Prior_Allowed") or "",
            notes=row.get("Notes") or "",
            context={name: row.get(name) for name in CONTEXT_FIELDS},
        )
        for row in data["Input_Meci"]
    )
    history = tuple(
        NativeHistoryRow(
            event_id=row["Event_ID"],
            league=row["League"],
            season=row["Season"],
            match_date=_parse_dt(row["Match_Date"]).date(),
            home=row["Home"],
            away=row["Away"],
            home_corners=int(row["HC"]),
            away_corners=int(row["AC"]),
            source_id=row["Source_ID"],
            source_url=row["Source_URL"],
            retrieved_utc=_parse_dt(row["Retrieved_UTC"]),
            definition=row["Definition"],
            evidence=row["Evidence"],
        )
        for row in data["Istoric_Nativ"]
    )
    sources = tuple(
        NativeSource(
            source_id=row["Source_ID"],
            league=row["League"],
            season=row["Season"],
            provider=row["Provider"],
            url=row["URL"],
            retrieved_utc=_parse_dt(row["Retrieved_UTC"]),
            coverage=row["Coverage"],
            status=row["Status"],
            rows=int(row["Rows"]),
            missing_corners=int(row["Missing_Corners"]),
            attempts=row["Attempts"],
            sha256=row["SHA256"],
        )
        for row in data["Surse_Import"]
    )
    return CornersBatchInput(matches=matches, history=history, sources=sources)


@lru_cache(maxsize=1)
def _result():
    return compute_corners_batch(_batch())


def _coerce(text: str):
    """Transformă o celulă CSV în valoarea Excel echivalentă."""
    if text == "":
        return None
    try:
        return float(text)
    except ValueError:
        return text


def _same(actual, expected, *, column: str) -> bool:
    if expected is None:
        return actual is None or actual == ""
    if isinstance(expected, float):
        if isinstance(actual, bool) or not isinstance(actual, (int, float)):
            return False
        return abs(float(actual) - expected) <= max(ABS_TOL, REL_TOL * abs(expected))
    text = str(actual) if actual is not None else ""
    if column in DASH_NORMALIZED:
        text = text.replace("\u2013", "-").replace("\u2014", "-")
        expected = expected.replace("\u2013", "-").replace("\u2014", "-")
    return text == expected


def test_workbook_sha256_matches_source_manifest():
    """Rezultatele așteptate sunt valabile doar pentru acest blob de șablon."""
    manifest = json.loads((PACK / "manifest_sursa.json").read_text(encoding="utf-8"))
    digest = sha256(TEMPLATE.read_bytes()).hexdigest()
    declared = [
        str(entry.get("sha256"))
        for entry in _iter_manifest_entries(manifest)
        if str(entry.get("sha256"))
    ]
    assert declared, "manifest_sursa.json nu declară niciun sha256"
    assert digest in declared, f"sha256 șablon {digest} nu apare în manifest"


def _iter_manifest_entries(manifest):
    if isinstance(manifest, dict):
        if "sha256" in manifest:
            yield manifest
        for value in manifest.values():
            yield from _iter_manifest_entries(value)
    elif isinstance(manifest, list):
        for item in manifest:
            yield from _iter_manifest_entries(item)


def test_template_line_definitions_match_engine():
    """Cele 14 linii și ordinea lor sunt citite din șablon, nu presupuse."""
    workbook = load_corners_workbook()
    sheet = workbook["Analiza_Linii"]
    for line in CORNERS_LINES:
        row = 6 + line.offset
        assert sheet.cell(row=row, column=6).value == line.label
        assert sheet.cell(row=row, column=7).value == line.market_type
        assert int(sheet.cell(row=row, column=8).value) == line.k
    config = workbook["Calibrare_Linii"]
    for line in CORNERS_LINES:
        assert config.cell(row=6 + line.offset, column=1).value == line.label


def test_template_line_config_matches_fixture():
    """`Calibrare_Linii` nu este input al aplicației: șablonul trebuie să o conțină deja."""
    workbook = load_corners_workbook()
    sheet = workbook["Calibrare_Linii"]
    for row in _fixture()["Calibrare_Linii"]:
        excel_row = int(row["excel_row"])
        for column, name in enumerate(
            (
                "Line",
                "Type",
                "k",
                "Calibration_Factor",
                "PASS_Max",
                "WATCH_Max",
                "Review_Status",
                "Scope_League",
            ),
            start=1,
        ):
            actual = sheet.cell(row=excel_row, column=column).value
            expected = row[name]
            if isinstance(expected, (int, float)) and not isinstance(expected, bool):
                assert float(actual) == float(expected), name
            else:
                assert actual == expected, name


def test_fixture_has_three_matches_and_42_selections():
    assert len(_batch().matches) == 3
    assert len(_batch().history) == 380
    assert len(_expected()) == 42


@pytest.mark.parametrize("index", range(42))
def test_expected_selection_matches_engine(index: int):
    expected_row = _expected()[index]
    result = _result()
    match = result.matches[expected_row["Match_ID"]]
    line_offset = int(expected_row["excel_row"]) - 6 - (match.slot - 1) * 14
    item = match.lines[line_offset]

    assert item.line.label == expected_row["Line"]
    mismatches: list[str] = []
    for column, raw in expected_row.items():
        if column == "excel_row":
            continue
        expected = _coerce(raw)
        actual = item.cells.get(column)
        if not _same(actual, expected, column=column):
            mismatches.append(f"{column}: actual={actual!r} expected={expected!r}")
    assert not mismatches, "\n".join(mismatches)


def test_replay_matches_are_excluded_from_top_live():
    """Exemplele REPLAY nu intră în clasamentul LIVE, nici prin selectorul Python."""
    result = _result()
    assert rank_top_live(result.ordered()) == []
    for match in result.ordered():
        for item in match.lines:
            assert item.live_eligible_best is False


def test_single_batch_ranking_matches_workbook():
    """Selectorul Python reproduce `Rank_LIVE` calculat de workbook în același lot."""
    result = _result()
    workbook_ranks = {
        (match.match_id, item.line.label): item.rank_live
        for match in result.ordered()
        for item in match.lines
        if item.live_eligible_best
    }
    ordered = rank_top_live(result.ordered())
    python_ranks = {
        (match.match_id, item.line.label): item.rank_live
        for match in result.ordered()
        for item in match.lines
        if item in ordered
    }
    assert python_ranks == workbook_ranks


def test_release_and_g0_states_are_model_states_not_errors():
    result = _result()
    for match in result.ordered():
        assert match.release in {"READY", "NOT RELEASED"}
        assert match.g0 in {"PASS", "FAIL", "NOT RELEASED"}
        for item in match.lines:
            assert not (isinstance(item.verdict, str) and item.verdict.startswith("ERROR"))


def test_distribution_branch_is_reported_per_match():
    """Alegerea Poisson/NB este un intermediar auditabil, nu o presupunere."""
    result = _result()
    kinds = {match.distribution_kind for match in result.ordered()}
    # Cele trei exemple acoperă ambele ramuri, cu etichetele exacte din workbook.
    assert kinds == {"POISSON", "NEGATIVE BINOMIAL"}
    for match in result.ordered():
        size = match.motor.get("NB_Size")
        if match.distribution_kind == "NEGATIVE BINOMIAL":
            assert isinstance(size, (int, float)) and size > 0
        else:
            assert size in (None, "") or isinstance(size, (int, float))


def test_recent_window_thresholds_come_from_filter_cells():
    """Ferestrele recente sunt construite cu `_xlfn._xlws.FILTER` + `LARGE`.

    Celulele sunt evaluate direct, nu deduse din rezultatul final: dacă evaluatorul
    ar pierde suportul pentru matrici dinamice, pragurile de dată ar dispărea în
    liniște, iar `Core_Nativ` ar raporta `NO HISTORY` în loc de `DERIVED`.
    """
    workbook = load_corners_workbook()
    sheet = workbook["Core_Nativ"]
    cached = _cached_workbook()["Core_Nativ"]
    filter_rows = [
        row
        for row in range(6, 42)
        if isinstance(sheet[f"L{row}"].value, str)
        and "_xlfn._xlws.FILTER" in sheet[f"L{row}"].value
    ]
    assert filter_rows, "Nicio celulă `Core_Nativ!L` nu mai folosește FILTER"

    extras = compute_corners_batch(
        _batch(), extra_cells=[f"Core_Nativ!L{row}" for row in filter_rows]
    ).extras
    windows = 0
    for row in filter_rows:
        expected = cached[f"L{row}"].value
        actual = extras[f"Core_Nativ!L{row}"]
        if isinstance(expected, datetime):
            windows += 1
            assert isinstance(actual, (int, float)), f"L{row}: {actual!r}"
            assert _serial_to_datetime(actual) == expected, f"L{row}"
        else:
            assert actual in (False, None, "", expected), f"L{row}: {actual!r}"
    assert windows >= 1, "Exemplele nu conțin nicio fereastră RECENT calculată"


def test_native_core_rows_are_derived_not_empty():
    """`Data_Status` confirmă că istoricul nativ a fost efectiv consumat."""
    statuses = {
        row.get("Data_Status")
        for match in _result().ordered()
        for row in match.core
    }
    assert "DERIVED" in statuses
    assert "NO HISTORY" not in statuses


def test_mu_and_d_are_consistent_across_lines():
    """Un singur calcul comun per meci: Mu și D sunt identice pe toate cele 14 linii."""
    for match in _result().ordered():
        mus = {item.mu for item in match.lines}
        ds = {item.dispersion for item in match.lines}
        assert len(mus) == 1
        assert len(ds) == 1
        assert math.isclose(
            next(iter(mus)), float(match.motor["Mu"]), rel_tol=REL_TOL, abs_tol=ABS_TOL
        )
