"""Oracle Microsoft Excel pentru Șansă Dublă V2.

Implicit SKIPPED: `pytest -q` nu pornește Excel. Pentru certificare:

    $env:PONTIFYBET_EXCEL_ORACLE = "1"
    python -m pytest tests/test_double_chance_excel_oracle.py -q

sau scriptul oficial `docs/PontifyBet_Sansa_Dubla_V2_Pachet_Cursor/verifica_excel.ps1`.
"""

from __future__ import annotations

import json
import math
import os
import time
from pathlib import Path

import pytest

from src.engines.double_chance import compute_double_chance

ROOT = Path(__file__).resolve().parent.parent
SPEC = ROOT / "docs" / "PontifyBet_Sansa_Dubla_V2_Pachet_Cursor"
TEMPLATE = ROOT / "templates" / "11_model_analiza_pariu_sansadubla_optimizat_V2.xlsx"
TOL = 1e-12
ORACLE_ENABLED = os.environ.get("PONTIFYBET_EXCEL_ORACLE") == "1"


def _blank(value) -> bool:
    return value is None or value == ""


def _same(actual, expected, *, exact_number: bool = False) -> bool:
    if _blank(actual) and _blank(expected):
        return True
    if isinstance(actual, (int, float)) and not isinstance(actual, bool):
        if isinstance(expected, (int, float)) and not isinstance(expected, bool):
            if exact_number:
                return float(actual) == float(expected)
            return math.isclose(float(actual), float(expected), rel_tol=TOL, abs_tol=TOL)
        return False
    return actual == expected


def _recalculate_with_excel(path: Path) -> None:
    import pythoncom
    import win32com.client

    pythoncom.CoInitialize()
    excel = win32com.client.DispatchEx("Excel.Application")
    excel.Visible = False
    excel.DisplayAlerts = False
    excel.AskToUpdateLinks = False
    excel.EnableEvents = False
    wb = None
    try:
        wb = excel.Workbooks.Open(str(path), UpdateLinks=0, ReadOnly=False)
        excel.Calculation = -4105
        excel.CalculateFullRebuild()
        deadline = time.time() + 90
        while int(excel.CalculationState) != 0 and time.time() < deadline:
            time.sleep(0.2)
        if int(excel.CalculationState) != 0:
            raise RuntimeError("Excel CalculationState nu a ajuns xlDone")
        wb.Save()
        wb.Close(SaveChanges=True)
        wb = None
    finally:
        if wb is not None:
            try:
                wb.Close(SaveChanges=False)
            except Exception:
                pass
        try:
            excel.Quit()
        except Exception:
            pass
        try:
            pythoncom.CoUninitialize()
        except Exception:
            pass


def _exact_score_or_level(cell: str) -> bool:
    return bool(
        cell.startswith("Double_Chance!L") or cell.startswith("Double_Chance!M")
    ) and cell[-1] in "456"


@pytest.mark.skipif(
    not ORACLE_ENABLED,
    reason="Microsoft Excel oracle SKIPPED (setează PONTIFYBET_EXCEL_ORACLE=1)",
)
def test_microsoft_excel_oracle_official_cells(tmp_path: Path):
    from openpyxl import load_workbook

    suite = json.loads((SPEC / "cazuri_test.json").read_text(encoding="utf-8"))
    mismatches: list[str] = []

    for case in suite["cases"]:
        dest = tmp_path / f"{case['id']}.xlsx"
        dest.write_bytes(TEMPLATE.read_bytes())
        wb = load_workbook(dest)
        for addr, value in case["inputs"].items():
            sheet, coord = addr.split("!", 1)
            cell = wb[sheet][coord]
            if isinstance(cell.value, str) and cell.value.startswith("="):
                mismatches.append(f"{case['id']} {addr}: input peste formulă")
                continue
            if isinstance(value, str) and value.replace(".", "", 1).lstrip("+-").isdigit():
                cell.number_format = "@"
                cell.value = value
            else:
                cell.value = value
        wb.save(dest)
        wb.close()

        try:
            _recalculate_with_excel(dest)
        except Exception as exc:
            pytest.skip(f"Microsoft Excel indisponibil sau instabil: {exc}")

        cached = load_workbook(dest, data_only=True)
        python = compute_double_chance(case["inputs"], evaluate=case["expected"].keys())
        for addr, expected in case["expected"].items():
            sheet, coord = addr.split("!", 1)
            excel_value = cached[sheet][coord].value
            exact = _exact_score_or_level(addr)
            if not _same(python.cells.get(addr), excel_value, exact_number=exact):
                mismatches.append(
                    f"{case['id']} {addr}: python={python.cells.get(addr)!r} excel={excel_value!r} expected={expected!r}"
                )
        cached.close()

    assert not mismatches, "Diferențe față de Microsoft Excel:\n" + "\n".join(mismatches)
