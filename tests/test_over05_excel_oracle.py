"""Oracle Microsoft Excel pentru Over 0.5 V4. SKIPPED dacă Excel lipsește."""

from __future__ import annotations

import json
import math
import time
from pathlib import Path

import pytest

from src.engines.over05 import OFFICIAL_COLUMNS, compute_over05
from src.engines.over05.inputs import ANALYSIS_SHEET, INPUT_COLUMNS

ROOT = Path(__file__).resolve().parent.parent
SPEC = ROOT / "docs" / "Pontifybet_Over05_V4_specificatii"
TEMPLATE = ROOT / "templates" / "1_model_analiza_over0.5_optimizat_V4.xlsx"
TOL = 1e-12


def _excel_available() -> bool:
    try:
        import pythoncom
        import win32com.client
    except ImportError:
        return False
    excel = None
    try:
        pythoncom.CoInitialize()
        excel = win32com.client.DispatchEx("Excel.Application")
        excel.Visible = False
        excel.DisplayAlerts = False
        return True
    except Exception:
        return False
    finally:
        if excel is not None:
            try:
                excel.Quit()
            except Exception:
                pass
        try:
            pythoncom.CoUninitialize()
        except Exception:
            pass


def _blank(value) -> bool:
    return value is None or value == ""


def _same(actual, expected) -> bool:
    if _blank(actual) and _blank(expected):
        return True
    if isinstance(actual, (int, float)) and not isinstance(actual, bool):
        if isinstance(expected, (int, float)) and not isinstance(expected, bool):
            return math.isclose(float(actual), float(expected), rel_tol=TOL, abs_tol=TOL)
        return False
    return actual == expected


def _recalculate_with_excel(path: Path) -> None:
    """Deschide o copie, CalculateFullRebuild, salvează valorile native."""
    import pythoncom
    import win32com.client

    pythoncom.CoInitialize()
    excel = win32com.client.DispatchEx("Excel.Application")
    excel.Visible = False
    excel.DisplayAlerts = False
    excel.AskToUpdateLinks = False
    excel.EnableEvents = False
    wb = None
    try:
        wb = excel.Workbooks.Open(str(path), UpdateLinks=0, ReadOnly=False)
        excel.Calculation = -4105
        excel.CalculateFullRebuild()
        deadline = time.time() + 90
        while int(excel.CalculationState) != 0 and time.time() < deadline:
            time.sleep(0.2)
        if int(excel.CalculationState) != 0:
            raise RuntimeError("Excel CalculationState nu a ajuns xlDone")
        wb.Save()
        wb.Close(SaveChanges=True)
        wb = None
    finally:
        if wb is not None:
            try:
                wb.Close(SaveChanges=False)
            except Exception:
                pass
        try:
            excel.Quit()
        except Exception:
            pass
        try:
            pythoncom.CoUninitialize()
        except Exception:
            pass


@pytest.mark.skipif(not _excel_available(), reason="Microsoft Excel indisponibil — oracle SKIPPED")
def test_microsoft_excel_oracle_official_cells(tmp_path: Path):
    from openpyxl import load_workbook

    fixtures = json.loads((SPEC / "golden_fixtures.json").read_text(encoding="utf-8"))
    mismatches: list[str] = []

    for fixture in fixtures["fixtures"]:
        dest = tmp_path / f"{fixture['match_id']}.xlsx"
        dest.write_bytes(TEMPLATE.read_bytes())
        wb = load_workbook(dest)
        ws = wb[ANALYSIS_SHEET]
        row = int(fixture["excel_row"])
        for col in INPUT_COLUMNS:
            ws[f"{col}{row}"].value = fixture["inputs"].get(col)
        wb.save(dest)
        wb.close()

        _recalculate_with_excel(dest)
        cached = load_workbook(dest, data_only=True)
        ws = cached[ANALYSIS_SHEET]
        python = compute_over05(fixture["inputs"], excel_row=row)
        for col in OFFICIAL_COLUMNS:
            excel_value = ws[f"{col}{row}"].value
            if not _same(python.cells[col], excel_value):
                mismatches.append(
                    f"{fixture['match_id']} {col}: python={python.cells[col]!r} excel={excel_value!r}"
                )
        cached.close()

    assert not mismatches, "Diferențe față de Microsoft Excel:\n" + "\n".join(mismatches)
